"""Genera el modelo de costo de capital de Aeropuertos del Neuquén S.A. (AdeN)."""
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SALIDA = "caso_aeropuerto/Modelo_Costo_Capital_AdeN.xlsx"

AZUL, VERDE, NEGRO = "0000FF", "008000", "000000"
F = "Arial"
fill_amarillo = PatternFill("solid", fgColor="FFFF00")
fill_titulo = PatternFill("solid", fgColor="1F3864")
fill_seccion = PatternFill("solid", fgColor="D9E1F2")
fill_reco = PatternFill("solid", fgColor="E2EFDA")
borde = Border(bottom=Side(style="thin", color="BFBFBF"))
PCT, PCT2 = "0.0%;(0.0%);-", "0.00%;(0.00%);-"
NUM, NUM1, NUM2 = "#,##0;(#,##0);-", "#,##0.0;(#,##0.0);-", "0.00"

wb = Workbook()


def titulo(ws, texto, sub, ancho="H"):
    ws.merge_cells(f"A1:{ancho}1")
    ws["A1"] = texto
    ws["A1"].font = Font(name=F, size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = fill_titulo
    ws.row_dimensions[1].height = 24
    ws["A2"] = sub
    ws["A2"].font = Font(name=F, size=9, italic=True, color="595959")


def seccion(ws, celda, texto, hasta):
    ws[celda] = texto
    fila = ws[celda].row
    for c in ws[f"A{fila}:{hasta}{fila}"][0]:
        c.fill = fill_seccion
        c.font = Font(name=F, bold=True)


def put(ws, celda, valor, color=NEGRO, fmt=None, bold=False, nota=None):
    ws[celda] = valor
    ws[celda].font = Font(name=F, color=color, bold=bold)
    if fmt:
        ws[celda].number_format = fmt
    if nota:
        ws[celda].comment = Comment(nota, "Modelo")


def txt(ws, celda, valor, bold=False, italic=False, color=NEGRO, wrap=False):
    ws[celda] = valor
    ws[celda].font = Font(name=F, bold=bold, italic=italic, color=color)
    if wrap:
        ws[celda].alignment = Alignment(wrap_text=True, vertical="top")


# ---------------------------------------------------------------- Datos
d = wb.active
d.title = "Datos"
titulo(d, "Datos del caso – Aeropuertos del Neuquén S.A. (A)",
       "Insumos tomados del caso IAE FZ-C-204-IA-1-s. Azul = dato ingresado; negro = fórmula.", "F")
for col, w in zip("ABCDEF", (52, 14, 14, 14, 14, 40)):
    d.column_dimensions[col].width = w

seccion(d, "A4", "1. Tasa libre de riesgo – Bonos del Tesoro de EE.UU. (Anexo 2)", "F")
rf = [("Bono 5 años – spot 23/12/2004", 0.0358),
      ("Bono 10 años – spot 23/12/2004", 0.0422),
      ("Bono 30 años – spot 23/12/2004", 0.0484),
      ("Bono 10 años – promedio ene-2000 a dic-2004", 0.0477)]
for i, (k, v) in enumerate(rf, start=5):
    txt(d, f"A{i}", k)
    put(d, f"B{i}", v, AZUL, PCT2)
txt(d, "F7", "Fuente: Anexo 2 y p. 5 del caso", italic=True, color="595959")

seccion(d, "A10", "2. Prima de riesgo de mercado EE.UU. hasta 2003 (Anexo 2)", "F")
for c, h in zip("ABCDE", ("Período", "Años", "Ibbotson (arit.)", "Damodaran (arit.)", "Damodaran (geom.)")):
    txt(d, f"{c}11", h, bold=True)
erp = [("1926/28 – 2003", 77, .072, .065, .048), ("1933 – 2003", 70, .081, .078, .059),
       ("1945 – 2003", 58, .074, .074, .058), ("1955 – 2003", 48, .056, .057, .041),
       ("1965 – 2003", 38, .043, .041, .029), ("1975 – 2003", 28, .070, .058, .040),
       ("1990 – 2003", 13, .060, .041, .036)]
for i, fila in enumerate(erp, start=12):
    txt(d, f"A{i}", fila[0])
    put(d, f"B{i}", fila[1], AZUL, "0")
    for c, v in zip("CDE", fila[2:]):
        put(d, f"{c}{i}", v, AZUL, PCT)
txt(d, "A19", "Promedio de todas las estimaciones", italic=True)
put(d, "C19", "=AVERAGE(C12:C18)", fmt=PCT)
put(d, "D19", "=AVERAGE(D12:D18)", fmt=PCT)
put(d, "E19", "=AVERAGE(E12:E18)", fmt=PCT)
for c, h in zip("BCD", ("Mínimo", "Máximo", "Punto medio")):
    txt(d, f"{c}21", h, bold=True)
txt(d, "A22", "ERP en decisiones regulatorias: EE.UU., Australia, Irlanda")
put(d, "B22", 0.05, AZUL, PCT); put(d, "C22", 0.07, AZUL, PCT); put(d, "D22", "=AVERAGE(B22:C22)", fmt=PCT)
txt(d, "A23", "ERP en decisiones regulatorias: Reino Unido")
put(d, "B23", 0.025, AZUL, PCT); put(d, "C23", 0.05, AZUL, PCT); put(d, "D23", "=AVERAGE(B23:C23)", fmt=PCT)
txt(d, "F22", "Fuente: p. 5 del caso", italic=True, color="595959")

seccion(d, "A25", "3. Betas de activos de operadores de aeropuertos (Anexo 3)", "F")
txt(d, "A26", "Empresa", bold=True); txt(d, "B26", "Beta activos", bold=True); txt(d, "C26", "Regulación", bold=True)
betas = [("Sydney Airport – SACL (2000)", 0.70), ("Auckland Intl. Airport – AIAL (2001)", 0.55),
         ("Melbourne Airport – MUDT (2000)", 0.75), ("Dublin Airport Authority – DAA (2005)", 0.61)]
for i, (k, v) in enumerate(betas, start=27):
    txt(d, f"A{i}", k)
    put(d, f"B{i}", v, AZUL, NUM2)
    txt(d, f"C{i}", "Price cap")
txt(d, "A31", "Promedio comparables", bold=True)
put(d, "B31", "=AVERAGE(B27:B30)", fmt="0.000", bold=True)
txt(d, "A32", "Mediana comparables")
put(d, "B32", "=MEDIAN(B27:B30)", fmt="0.000")
txt(d, "F27", "Los cuatro comparables operan bajo price cap, como AdeN.", italic=True, color="595959", wrap=True)

seccion(d, "A34", "4. Ajustes de riesgo (pp. 6–9 del caso)", "F")
txt(d, "A35", "Prima por riesgo regulatorio (consenso de mercado)")
put(d, "B35", 0.02, AZUL, PCT, nota="p. 7: 'cierto consenso en utilizar un 2%'. Se calcula como (βGB − βUSA) × ERP para llevar betas de EE.UU. (rate of return) a price cap.")
txt(d, "A36", "Ajuste a la beta por menor calificación (precedente DAA)")
put(d, "B36", 0.20, AZUL, PCT, nota="p. 8: a la DAA se le subió la beta 20% cuando S&P la bajó de A+ a A. AA2000 tiene Baa1/BBB+.")
txt(d, "A37", "Ajuste a la beta por tamaño (supuesto, igual al de la DAA)")
put(d, "B37", 0.20, AZUL, PCT, nota="Supuesto del modelo: el caso no da un factor. Se replica el 20% de la DAA para representar la posición máxima de la empresa.")

seccion(d, "A39", "5. Riesgo país – EMBI+ en puntos básicos (Anexo 3)", "F")
txt(d, "A40", "Medida", bold=True); txt(d, "B40", "Argentina", bold=True); txt(d, "C40", "EMBI+ global", bold=True)
rp = [("Spot 23/12/2004", 4668, 354), ("Promedio diciembre 2004", 5027, 368),
      ("Promedio año 2004", 5210, 437), ("Promedio 1996 – 2004", 2857, 752),
      ("Promedio año 2003", 5561, 561), ("Promedio 1996 – 2000 (pre-crisis)", 659, 887)]
for i, (k, a, g) in enumerate(rp, start=41):
    txt(d, f"A{i}", k)
    put(d, f"B{i}", a, AZUL, NUM)
    put(d, f"C{i}", g, AZUL, NUM)
txt(d, "A47", "Promedio latinoamericano usado por la CREG de Colombia (2003)")
put(d, "B47", 619, AZUL, NUM)
txt(d, "A48", "EPRE San Juan (2003): no incluyó riesgo país")
put(d, "B48", 0, AZUL, NUM)

seccion(d, "A50", "6. AdeN – datos operativos y contables 2004 (Anexos 1 y 5)", "F")
aden = [("Base de activos regulados – Total inversión permanente (miles $)", 5729, NUM,
         "Anexo 5, 2004. Se usa como aproximación de la RAB."),
        ("Bienes de uso (miles $)", 7967.4, NUM1, "Anexo 1, 2004"),
        ("Pasajeros 2004 (miles)", 320.18, NUM1, "Anexo 5"),
        ("Ingresos totales 2004 (miles $)", 1994.6, NUM1, "Anexo 1"),
        ("Ingresos aeronáuticos 2004 (miles $)", 1373.0, NUM1, "Anexo 1")]
for i, (k, v, fm, nota) in enumerate(aden, start=51):
    txt(d, f"A{i}", k)
    put(d, f"B{i}", v, AZUL, fm, nota=nota)
txt(d, "A56", "Ingreso total por pasajero ($)")
put(d, "B56", "=B54/B53", fmt=NUM2)
txt(d, "A57", "Ingreso aeronáutico por pasajero ($)")
put(d, "B57", "=B55/B53", fmt=NUM2)

seccion(d, "A59", "7. Estructura de capital e impuestos", "F")
txt(d, "A60", "Deuda financiera / (D + E)")
put(d, "B60", 0, AZUL, PCT, nota="p. 4: las partes acordaron tratar los pasivos como espontáneos, así que no hay deuda financiera. Por eso WACC = Ke.")
txt(d, "A61", "Costo de la deuda Kd (no aplica)")
put(d, "B61", 0, AZUL, PCT)
txt(d, "A62", "Tasa de impuesto a las ganancias T")
put(d, "B62", 0.35, AZUL, PCT, nota="Supuesto: tasa general de Argentina en 2004 (35%). El caso no la informa. No afecta el WACC porque D = 0.")
txt(d, "A63", "Tipo de cambio de referencia ($ por US$)")
put(d, "B63", 3, AZUL, "0.00", nota="Anexo 1, nota (**): estable cerca de $3 por dólar desde 2002.")

# ---------------------------------------------------------------- Escenarios
e = wb.create_sheet("Alternativas")
titulo(e, "Costo de capital de AdeN – tres alternativas de negociación",
       "Ke = rf + β·ERP + prima regulatoria + riesgo país. En dólares nominales. Verde = vínculo a 'Datos'; celdas amarillas = palancas de cada alternativa.", "E")
for col, w in zip("ABCDE", (50, 20, 20, 20, 4)):
    e.column_dimensions[col].width = w
heads = ["A. Postura del regulador", "B. Máximo de la empresa", "C. Propuesta recomendada"]
txt(e, "A4", "Componente", bold=True)
for c, h in zip("BCD", heads):
    txt(e, f"{c}4", h, bold=True, wrap=True)
    e[f"{c}4"].alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
e.row_dimensions[4].height = 30
e["D4"].fill = fill_reco

filas = {
    5: ("Tasa libre de riesgo rf", ["=Datos!B6", "=Datos!B7", "=Datos!B8"], PCT2),
    6: ("Prima de riesgo de mercado ERP", ["=Datos!E12", "=Datos!C12", "=Datos!D22"], PCT),
    7: ("Beta de activos (promedio comparables)", ["=Datos!B31", "=Datos!B31", "=Datos!B31"], "0.000"),
    8: ("Ajuste por calificación crediticia", [0, "=Datos!B36", "=Datos!B36"], PCT),
    9: ("Ajuste por tamaño (un solo aeropuerto regional)", [0, "=Datos!B37", 0], PCT),
    11: ("Prima por riesgo regulatorio", [0, "=Datos!B35", 0], PCT),
    12: ("Riesgo país (puntos básicos)", ["=Datos!B48", "=Datos!B44", "=Datos!B47"], NUM),
}
for r, (lab, vals, fm) in filas.items():
    txt(e, f"A{r}", lab)
    for c, v in zip("BCD", vals):
        color = VERDE if isinstance(v, str) else AZUL
        put(e, f"{c}{r}", v, color, fm)
        e[f"{c}{r}"].fill = fill_amarillo
txt(e, "A10", "Beta apalancada ajustada (D = 0 ⇒ βE = βA)", bold=True)
txt(e, "A13", "Riesgo país (%)")
txt(e, "A15", "Costo del capital propio Ke", bold=True)
txt(e, "A16", "Peso de la deuda D/(D+E)")
txt(e, "A17", "Kd × (1 − T)")
txt(e, "A18", "WACC (tasa de descuento)", bold=True)
for c in "BCD":
    put(e, f"{c}10", f"={c}7*(1+{c}8)*(1+{c}9)", fmt="0.000", bold=True)
    put(e, f"{c}13", f"={c}12/10000", fmt=PCT2)
    put(e, f"{c}15", f"={c}5+{c}10*{c}6+{c}11+{c}13", fmt=PCT2, bold=True)
    put(e, f"{c}16", "=Datos!B60", VERDE, PCT)
    put(e, f"{c}17", "=Datos!B61*(1-Datos!B62)", VERDE, PCT2)
    put(e, f"{c}18", f"={c}17*{c}16+{c}15*(1-{c}16)", fmt=PCT2, bold=True)
    e[f"{c}18"].border = Border(top=Side(style="thin"), bottom=Side(style="double"))

seccion(e, "A20", "Impacto sobre la tarifa (costo anual del capital = WACC × RAB)", "D")
txt(e, "A21", "Base de activos regulados – RAB (miles $)")
txt(e, "A22", "Costo anual del capital (miles $)", bold=True)
txt(e, "A23", "Pasajeros 2004 (miles)")
txt(e, "A24", "Costo del capital por pasajero ($)", bold=True)
txt(e, "A25", "Ingreso total por pasajero 2004 ($)")
txt(e, "A26", "Costo del capital / ingresos totales 2004")
txt(e, "A27", "Diferencia de WACC frente a la propuesta C (pp)")
txt(e, "A28", "Valor presente de $100 anuales a perpetuidad ($)")
for c in "BCD":
    put(e, f"{c}21", "=Datos!B51", VERDE, NUM)
    put(e, f"{c}22", f"={c}18*{c}21", fmt=NUM, bold=True)
    put(e, f"{c}23", "=Datos!B53", VERDE, NUM1)
    put(e, f"{c}24", f"=IF({c}23=0,0,{c}22/{c}23)", fmt=NUM2, bold=True)
    put(e, f"{c}25", "=Datos!B56", VERDE, NUM2)
    put(e, f"{c}26", f"=IF(Datos!$B$54=0,0,{c}22/Datos!$B$54)", fmt=PCT)
    put(e, f"{c}27", f"=({c}18-$D$18)*100", fmt="0.0;(0.0);-")
    put(e, f"{c}28", f"=IF({c}18=0,0,100/{c}18)", fmt=NUM)

seccion(e, "A30", "Criterio de cada alternativa", "D")
crit = [
    ("rf", "Bono a 10 años del día (4,22%), el valor más bajo y actual.",
     "Bono a 30 años (4,84%): coincide con un plazo de concesión largo.",
     "Promedio del bono a 10 años 2000–2004 (4,77%): suaviza un mínimo coyuntural."),
    ("ERP", "Damodaran, media geométrica 1928–2003 (4,8%).",
     "Ibbotson, media aritmética 1926–2003 (7,2%).",
     "Punto medio de lo aceptado por reguladores con price cap en EE.UU., Australia e Irlanda (6,0%)."),
    ("β", "Promedio de comparables, sin ajustes.",
     "+20% por calificación y +20% por tamaño.",
     "+20% por calificación (precedente DAA: A+→A; AA2000 es BBB+)."),
    ("Prima regulatoria", "No aplica: los comparables ya están regulados por price cap.",
     "Agrega el 2% de 'consenso'.",
     "No se suma: los 4 comparables ya son price cap; sumarla es contar el riesgo dos veces."),
    ("Riesgo país", "Cero (precedente EPRE San Juan; estudio McKinsey).",
     "Promedio Argentina 1996–2004 (2.857 pb). Con el spot (4.668 pb) sería aún mayor.",
     "Riesgo normalizado: 619 pb (CREG Colombia). Es casi igual al de Argentina antes de la crisis (659 pb)."),
]
for i, fila in enumerate(crit, start=31):
    txt(e, f"A{i}", fila[0], bold=True)
    for c, t in zip("BCD", fila[1:]):
        txt(e, f"{c}{i}", t, wrap=True)
    e.row_dimensions[i].height = 62

seccion(e, "A37", "Evaluación", "D")
ev = [("¿La acepta la contraparte?", "Sí (es su propuesta)", "No: duplica la tarifa", "Probable: cada supuesto tiene precedente"),
      ("¿Atrae capital?", "No: menos que un bono argentino", "Sí, en exceso", "Sí: cubre el riesgo argentino"),
      ("¿Sostiene el negocio?", "No: consolida las pérdidas", "No: rompe la negociación", "Sí")]
for i, fila in enumerate(ev, start=38):
    for c, t in zip("ABCD", fila):
        txt(e, f"{c}{i}", t, bold=(c == "A"), wrap=True)
    e.row_dimensions[i].height = 30
for r in range(31, 41):
    e[f"D{r}"].fill = fill_reco

graf = BarChart()
graf.type = "col"
graf.title = "WACC por alternativa"
graf.y_axis.title = "WACC"
graf.y_axis.numFmt = "0%"
graf.add_data(Reference(e, min_col=2, max_col=4, min_row=18), from_rows=True, titles_from_data=False)
graf.set_categories(Reference(e, min_col=2, max_col=4, min_row=4))
graf.legend = None
graf.dataLabels = DataLabelList()
graf.dataLabels.showVal = True
graf.height, graf.width = 7.5, 15
e.add_chart(graf, "F4")

# ---------------------------------------------------------------- Sensibilidad
s = wb.create_sheet("Sensibilidad")
titulo(s, "Sensibilidad del Ke de la propuesta C",
       "Mantiene rf, β ajustada y prima regulatoria de la alternativa C (hoja 'Alternativas') y varía ERP, riesgo país y β.", "H")
s.column_dimensions["A"].width = 24
for c in "BCDEFGH":
    s.column_dimensions[c].width = 13
seccion(s, "A4", "Tabla 1. Ke según ERP (filas) y riesgo país en pb (columnas)", "H")
txt(s, "A5", "ERP \\ Riesgo país", bold=True)
rps = [0, 354, 619, 659, 1500, 2857, 4668]
for c, v in zip("BCDEFGH", rps):
    put(s, f"{c}5", v, AZUL, NUM, bold=True)
erps = [0.029, 0.041, 0.048, 0.055, 0.060, 0.065, 0.072, 0.081]
for i, v in enumerate(erps, start=6):
    put(s, f"A{i}", v, AZUL, PCT, bold=True)
    for c in "BCDEFGH":
        put(s, f"{c}{i}", f"=Alternativas!$D$5+Alternativas!$D$10*$A{i}+Alternativas!$D$11+{c}$5/10000", fmt=PCT)
txt(s, "A14", "Fila resaltada: ERP de la propuesta (6,0%). Columna resaltada: riesgo país normalizado (619 pb).", italic=True, color="595959")
for c in "BCDEFGH":
    s[f"{c}10"].fill = fill_reco
for i in range(6, 14):
    s[f"D{i}"].fill = fill_reco

seccion(s, "A16", "Tabla 2. Ke según β ajustada (filas) y ERP (columnas), con riesgo país de 619 pb", "H")
txt(s, "A17", "β \\ ERP", bold=True)
erps2 = [0.041, 0.048, 0.055, 0.060, 0.065, 0.072]
for c, v in zip("BCDEFG", erps2):
    put(s, f"{c}17", v, AZUL, PCT, bold=True)
bs = [0.55, 0.6525, 0.70, 0.75, 0.783, 0.85, 0.94]
for i, v in enumerate(bs, start=18):
    put(s, f"A{i}", v, AZUL, "0.000", bold=True)
    for c in "BCDEFG":
        put(s, f"{c}{i}", f"=Alternativas!$D$5+$A{i}*{c}$17+Alternativas!$D$11+Alternativas!$D$13", fmt=PCT)

seccion(s, "A26", "Tabla 3. Costo del capital por pasajero ($) según riesgo país (ERP y β de la propuesta)", "H")
txt(s, "A27", "Riesgo país (pb)", bold=True)
txt(s, "A28", "WACC", bold=True)
txt(s, "A29", "$ por pasajero", bold=True)
txt(s, "A30", "% del ingreso total 2004", bold=True)
for c, v in zip("BCDEFGH", rps):
    put(s, f"{c}27", v, AZUL, NUM, bold=True)
    put(s, f"{c}28", f"=Alternativas!$D$5+Alternativas!$D$10*Alternativas!$D$6+Alternativas!$D$11+{c}27/10000", fmt=PCT)
    put(s, f"{c}29", f"={c}28*Datos!$B$51/Datos!$B$53", fmt=NUM2)
    put(s, f"{c}30", f"={c}28*Datos!$B$51/Datos!$B$54", fmt=PCT)

# ---------------------------------------------------------------- Resumen
r = wb.create_sheet("Resumen", 0)
titulo(r, "Resumen para la Junta Directiva – Tasa de descuento de AdeN",
       "Todos los valores se calculan en las hojas 'Alternativas' y 'Datos'.", "E")
for col, w in zip("ABCDE", (40, 18, 18, 18, 4)):
    r.column_dimensions[col].width = w
txt(r, "A4", "Indicador", bold=True)
for c, h in zip("BCD", heads):
    txt(r, f"{c}4", h, bold=True, wrap=True)
    r[f"{c}4"].alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
r.row_dimensions[4].height = 30
r["D4"].fill = fill_reco
res = [("WACC = Ke (US$ nominal)", 18, PCT2), ("Beta ajustada", 10, "0.000"),
       ("Riesgo país (pb)", 12, NUM), ("Costo anual del capital (miles $)", 22, NUM),
       ("Costo del capital por pasajero ($)", 24, NUM2), ("Costo del capital / ingresos 2004", 26, PCT)]
for i, (lab, fila, fm) in enumerate(res, start=5):
    txt(r, f"A{i}", lab)
    for c in "BCD":
        put(r, f"{c}{i}", f"=Alternativas!{c}{fila}", VERDE, fm)
        r[f"{c}{i}"].border = borde
    r[f"D{i}"].fill = fill_reco
txt(r, "A12", "Recomendación", bold=True)
r.merge_cells("A13:D15")
txt(r, "A13", "Proponer al regulador un WACC de alrededor de 15,7% en dólares nominales (alternativa C). "
    "Mínimo aceptable 12,1% (C con ERP 4,8% y riesgo país igual al EMBI+ global de 354 pb, ver 'Sensibilidad'). "
    "Abrir la negociación sumando la prima regulatoria del 2% (≈17,7%). "
    "Exigir que los flujos se expresen en dólares, o que se sume a la tasa el diferencial de inflación.", wrap=True)
r.row_dimensions[13].height = 20

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4" if ws.title != "Datos" else "A3"

wb.save(SALIDA)
print("ok", SALIDA)
