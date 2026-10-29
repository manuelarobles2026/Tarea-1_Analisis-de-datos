"""Informe de 2 páginas para la Junta Directiva, con las cifras del Excel del Grupo 3.

Cada cifra se lee de una celda del Excel del grupo (o se deriva de celdas, indicado en la
tabla de control). Uso: python generar_informe_jd.py ruta/Caso_2_Grupo_3_1.xlsx carpeta_salida
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import load_workbook

AZUL = "1F3864"
CELDAS = {
    "wacc": ("WACC_COMPARABLES", "K37"), "g": ("VALORACIÓN", "B5"), "ops": ("VALORACIÓN", "B8"),
    "tv_share": ("VALORACIÓN", "B9"), "noop": ("VALORACIÓN", "B16"), "claims": ("VALORACIÓN", "B25"),
    "eq": ("VALORACIÓN", "B26"), "sin_bh": ("VALORACIÓN", "B28"), "sin_rev": ("VALORACIÓN", "B29"),
    "eq_sin": ("VALORACIÓN", "B30"), "mult": ("VALORACIÓN", "B33"), "eq_mult": ("VALORACIÓN", "B36"),
    "g_impl": ("VALORACIÓN", "B37"), "ev_ebitda07": ("VALORACIÓN", "B38"), "eq_pe": ("VALORACIÓN", "B41"),
    "dv_ops": ("VALORACIÓN", "B42"), "apertura": ("VALORACIÓN", "B45"), "objetivo": ("VALORACIÓN", "B46"),
    "maximo": ("VALORACIÓN", "B47"), "afil": ("VALORACIÓN", "B13"), "afil_sh": ("VALORACIÓN", "B49"),
    "sinrev_sh": ("VALORACIÓN", "B50"), "prima": ("VALORACIÓN", "B51"), "mult_max": ("VALORACIÓN", "B52"),
    "m_ebitda": ("Exhibit 4", "J65"), "m_comp": ("WACC_COMPARABLES", "R13"), "rpm06": ("Exhibit 3", "F73"),
    "rpm07": ("Exhibit 3", "G73"), "cpa06": ("Exhibit 3", "F60"), "cpa07": ("Exhibit 3", "G60"),
    "deuda": ("Exhibit 6", "H8"), "linea11": ("PROYECCIONES", "F85"), "ahorro_bh12": ("SINERGIAS", "H8"),
    "ipo": ("PROYECCIONES", "K7"), "ebitda07": ("PROYECCIONES", "B26"),
    "dv_total": ("VALORACIÓN", "B43"), "ev_total": ("VALORACIÓN", "B24"), "afil_libros": ("Exhibit 5", "G72"),
    "linea12_total": ("PROYECCIONES", "G85"), "bullet": ("Exhibit 6", "P11"), "deuda_ebitda08": ("PROYECCIONES", "C97"),
    "deuda_ebitda12": ("PROYECCIONES", "G97"), "cobertura08": ("PROYECCIONES", "C98"),
}
DERIVADAS = {  # clave: (fórmula en celdas, función)
    "rpm_var": ("Exhibit 3!G73 / Exhibit 3!F73 − 1", lambda v: v["rpm07"] / v["rpm06"] - 1),
    "cpa_var": ("Exhibit 3!G60 / Exhibit 3!F60 − 1", lambda v: v["cpa07"] / v["cpa06"] - 1),
    "sin_tot": ("VALORACIÓN!B28 + VALORACIÓN!B29", lambda v: v["sin_bh"] + v["sin_rev"]),
    "mitad_rev": ("VALORACIÓN!B29 / 2", lambda v: v["sin_rev"] / 2),
    "linea12": ("PROYECCIONES!G85 − Exhibit 6!P11", lambda v: v["linea12_total"] - v["bullet"]),
    "deuda50": ("0,5 × VALORACIÓN!B8", lambda v: 0.5 * v["ops"]),
    "reduccion": ("Exhibit 6!H8 − 0,5 × VALORACIÓN!B8", lambda v: v["deuda"] - 0.5 * v["ops"]),
    "dv_afil_libros": ("Exhibit 6!H8 / (VALORACIÓN!B24 − VALORACIÓN!B13 + Exhibit 5!G72)",
                       lambda v: v["deuda"] / (v["ev_total"] - v["afil"] + v["afil_libros"])),
}


def n(x, d=0):
    s = f"{abs(x):,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("−" if x < 0 else "") + s


def p(x, d=1):
    return n(x * 100, d) + "%"


def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def bordes(celda, color="BFC9DC"):
    tcPr = celda._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), color)
        b.append(e)
    tcPr.append(b)


def margen_celda(celda, cm=0.12):
    tcPr = celda._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for lado in ("left", "right"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:w"), str(int(cm * 567)))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tcPr.append(mar)


def texto(par, partes, tam=10.5):
    """partes: lista de (texto, negrita)."""
    for t, b in partes:
        r = par.add_run(t)
        r.bold = b
        r.font.size = Pt(tam)
    return par


def main(ruta_xlsx, salida):
    wb = load_workbook(ruta_xlsx, data_only=True)
    v = {k: wb[h][c].value for k, (h, c) in CELDAS.items()}
    for k, val in v.items():
        assert isinstance(val, (int, float)), (k, val)
    for k, (_, f) in DERIVADAS.items():
        v[k] = f(v)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.5)
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.font.size = Pt(9.5)
    st.paragraph_format.space_after = Pt(2)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.line_spacing = 1.05

    def parrafo(partes=None, tam=10.5, despues=4, alin=None):
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(despues)
        if alin:
            par.alignment = alin
        if partes:
            texto(par, partes, tam)
        return par

    def titulo_seccion(t):
        par = parrafo(despues=1.5)
        par.paragraph_format.space_before = Pt(10)
        par.paragraph_format.keep_with_next = True
        r = par.add_run(t)
        r.bold = True
        r.font.size = Pt(12.5)
        r.font.color.rgb = RGBColor.from_string(AZUL)
        pPr = par._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bot = OxmlElement("w:bottom")
        bot.set(qn("w:val"), "single")
        bot.set(qn("w:sz"), "6")
        bot.set(qn("w:color"), AZUL)
        bdr.append(bot)
        pPr.append(bdr)

    def viñeta(partes, tam=10.5):
        par = doc.add_paragraph(style="List Bullet")
        par.paragraph_format.space_after = Pt(3)
        par.paragraph_format.left_indent = Cm(0.6)
        texto(par, partes, tam)
        return par

    def numerada(partes, tam=10.5):
        par = doc.add_paragraph(style="List Number")
        par.paragraph_format.space_after = Pt(4)
        par.paragraph_format.left_indent = Cm(0.6)
        texto(par, partes, tam)
        return par

    # Encabezado
    par = parrafo([("MEMORANDO A LA JUNTA DIRECTIVA · AMERICAN CABLE COMMUNICATIONS (ACC)", True)], tam=8, despues=0)
    par.runs[0].font.color.rgb = RGBColor.from_string("666666")
    par = parrafo([("De: Desarrollo de Negocios (Grupo 3) · Diciembre de 2007 · Todas las cifras en millones de US$", False)],
                  tam=8, despues=4)
    par.runs[0].font.color.rgb = RGBColor.from_string("666666")
    par = parrafo(despues=4)
    r = par.add_run(f"Compra de AirThread: recomendamos ofertar, con un precio máximo de US${n(v['maximo'])} millones")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor.from_string(AZUL)

    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    sombrear(c, "E8EEF7")
    bordes(c, AZUL)
    pr = c.paragraphs[0]
    pr.paragraph_format.space_after = Pt(3)
    pr.paragraph_format.space_before = Pt(3)
    texto(pr, [("Recomendación: ", True),
               ("comprar el 100% de AirThread Connections (ATC), siempre que el precio no supere "
                f"US${n(v['maximo'])} millones. Proponemos abrir la negociación en US${n(v['apertura'])} millones y buscar "
                f"cerrar cerca de US${n(v['objetivo'])} millones. Si el vendedor pide más, ACC debe retirarse: por encima de "
                "ese precio estaríamos pagando por beneficios inciertos que solo existirían gracias a ACC.", False)])
    parrafo(despues=0)

    titulo_seccion("1. La decisión que debe tomar la Junta")
    parrafo([("ACC ofrece televisión, internet y telefonía fija, pero no telefonía celular. Sus competidores ya venden "
              "paquetes que incluyen celular, y ACC corre el riesgo de quedarse atrás. Comprar ATC, un operador celular "
              "regional, resolvería ese vacío. La pregunta no es solo si comprar, sino ", False),
             ("cuánto pagar sin destruir valor para los accionistas", True),
             (": pagar de más es la forma más común en que una adquisición fracasa.", False)])

    titulo_seccion("2. Lo que sabemos")
    viñeta([("ATC es rentable, pero más débil que sus competidores. ", True),
            (f"De cada US$100 que vende, le quedan US${n(v['m_ebitda'] * 100, 0)} de utilidad operativa (antes de "
             f"depreciación), frente a US${n(v['m_comp'] * 100, 0)} de sus rivales. Además, en 2007 el precio que cobra por "
             f"minuto bajó {p(-v['rpm_var'])} y conseguir cada cliente nuevo le costó {p(v['cpa_var'])} más.", False)])
    viñeta([(f"Por sí sola, ATC vale unos US${n(v['eq'])} millones. ", True),
            ("Estimamos el dinero que generará en los próximos años y lo llevamos a valor de hoy, exigiendo un "
             f"rendimiento de {p(v['wacc'], 1)} anual, que es lo que cuesta financiar la compra. Ese valor equivale a "
             f"{n(v['ev_ebitda07'], 1)} veces su utilidad operativa de 2007.", False)])
    viñeta([(f"Junto con ACC, ATC vale US${n(v['sin_tot'])} millones más. ", True),
            (f"US${n(v['sin_bh'])} millones vienen de ahorrar costos: ATC podría usar la red de fibra de ACC para conectar "
             f"sus antenas, en lugar de alquilar líneas (US${n(v['ahorro_bh12'])} millones de ahorro al año en 2012). Los "
             f"otros US${n(v['sin_rev'])} millones vendrían de vender planes celulares a empresas clientes de ACC, pero "
             "esa cifra es una estimación sin experiencia previa que la respalde.", False)])
    viñeta([("El vendedor probablemente pedirá mucho más. ", True),
            (f"Si se usa el precio al que se negocian empresas parecidas, ATC valdría US${n(v['eq_mult'])} millones. "
             f"Para justificarlo, ATC tendría que crecer {p(v['g_impl'])} al año para siempre, algo poco realista para una "
             "empresa que está perdiendo margen.", False)])
    viñeta([("La financiación propuesta por el banco es exigente. ", True),
            (f"El banco ofrece prestar US${n(v['deuda'])} millones. Es el {p(v['dv_total'])} del valor total de ATC, "
             "justo por debajo del 50% que recomienda el banco. Sin embargo, el dinero que genera ATC no alcanza para pagar "
             f"las cuotas: necesitaría un crédito adicional que llegaría a US${n(v['linea12'])} millones en 2012, cuando "
             f"además habrá que renovar un saldo de US${n(v['bullet'])} millones.", False)])

    filas = [("ATC por sí sola", v["eq"], False),
             ("+ Ahorro por usar la red de fibra de ACC (más seguro)", v["sin_bh"], False),
             ("+ Ventas a empresas clientes de ACC (incierto)", v["sin_rev"], False),
             ("ATC con todos los beneficios de la unión", v["eq_sin"], True),
             ("Referencia: precio según empresas parecidas", v["eq_mult"], False)]
    t = doc.add_table(rows=len(filas) + 1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, (txt, der) in enumerate([("¿Cuánto vale ATC? (patrimonio, 100%)", False), ("US$ millones", True)]):
        cel = t.cell(0, j)
        sombrear(cel, AZUL)
        pr = cel.paragraphs[0]
        r = pr.add_run(txt)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        if der:
            pr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for i, (etq, val, neg) in enumerate(filas, start=1):
        for j, txt in enumerate([etq, n(val)]):
            cel = t.cell(i, j)
            bordes(cel)
            margen_celda(cel)
            if neg:
                sombrear(cel, "E8EEF7")
            pr = cel.paragraphs[0]
            pr.paragraph_format.space_after = Pt(0)
            r = pr.add_run(txt)
            r.bold = neg
            r.font.size = Pt(9.5)
            if i == len(filas):
                r.italic = True
            if j == 1:
                pr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for fila in t.rows:
        fila.cells[0].width = Cm(10.5)
        fila.cells[1].width = Cm(2.8)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(0)
    titulo_seccion("3. Alternativas")
    cols = ["", "A. Pagar solo lo que ATC vale por sí sola", "B. Compartir solo el ahorro de la fibra",
            "C. Pagar todos los beneficios de la unión", "D. No comprar y firmar un acuerdo comercial"]
    filas = [
        ("Precio", f"Hasta {n(v['eq'])}", f"Entre {n(v['apertura'])} y {n(v['maximo'])}", f"Hasta {n(v['eq_sin'])}",
         "No hay compra"),
        ("Qué gana ACC", f"Se queda con todos los beneficios de la unión ({n(v['sin_tot'])})",
         f"Se queda con las ventas a empresas ({n(v['sin_rev'])}) como colchón",
         "Nada: todo el beneficio se lo lleva el vendedor",
         f"Cobra a ATC por usar su fibra (hasta {n(v['ahorro_bh12'])} al año)"),
        ("Qué implica", "Es el precio más seguro, pero el vendedor difícilmente lo aceptará.",
         f"El vendedor recibe {p(v['prima'])} más de lo que ATC vale por sí sola.",
         f"Gana cualquier subasta, pero si las ventas a empresas salen a medias, ACC pierde {n(v['mitad_rev'])}.",
         "No requiere deuda, pero ACC sigue sin telefonía celular."),
        ("Mayor riesgo", "Que otro comprador se quede con ATC.", "Que el vendedor no acepte bajar su precio.",
         "Pagar de más.", "Que un competidor compre ATC."),
    ]
    t = doc.add_table(rows=len(filas) + 1, cols=5)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    anchos = [Cm(2.4), Cm(3.5), Cm(3.7), Cm(3.5), Cm(3.5)]
    for j, txt in enumerate(cols):
        cel = t.cell(0, j)
        sombrear(cel, "2E7D32" if j == 2 else AZUL)
        margen_celda(cel)
        pr = cel.paragraphs[0]
        pr.paragraph_format.space_after = Pt(0)
        r = pr.add_run(txt + (" (recomendada)" if j == 2 else ""))
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)
    for i, fila in enumerate(filas, start=1):
        for j, txt in enumerate(fila):
            cel = t.cell(i, j)
            bordes(cel)
            margen_celda(cel)
            if j == 2:
                sombrear(cel, "EAF4EA")
            elif j == 0:
                sombrear(cel, "F2F2F2")
            pr = cel.paragraphs[0]
            pr.paragraph_format.space_after = Pt(0)
            pr.paragraph_format.line_spacing = 1.0
            r = pr.add_run(txt)
            r.font.size = Pt(9)
            r.bold = j == 0
    for fila in t.rows:
        for j, w in enumerate(anchos):
            fila.cells[j].width = w

    titulo_seccion("4. Recomendación: alternativa B")
    numerada([("Es un precio razonable. ", True),
              (f"Pagar hasta US${n(v['maximo'])} millones significa reconocerle al vendedor {p(v['prima'])} sobre lo que ATC "
               f"vale por sí sola. Otro método, basado en las utilidades históricas del sector, da un valor parecido "
               f"(US${n(v['eq_pe'])} millones).", False)])
    numerada([("Solo pagamos lo que depende de nosotros. ", True),
              ("El ahorro de la fibra lo genera la red de ACC y es el beneficio más seguro. Las ventas a empresas "
               f"(US${n(v['sin_rev'])} millones) no se pagan por adelantado: quedan como colchón si algo sale mal.", False)])
    numerada([("Cómo negociar. ", True),
              (f"Abrir en US${n(v['apertura'])} millones: lo que ATC vale por sí sola menos el {p(v['ipo'], 0)} que le costaría "
               f"vender acciones en bolsa, que es la otra opción del vendedor. Buscar cerrar en US${n(v['objetivo'])} millones. "
               "Si el vendedor pide más, ofrecer un pago adicional solo si se cumplen las metas de ventas a empresas, en lugar "
               "de subir el precio fijo.", False)])
    numerada([("La financiación es viable si ACC la respalda. ", True),
              (f"El préstamo de US${n(v['deuda'])} millones está dentro del límite del banco, y la deuda baja de "
               f"{n(v['deuda_ebitda08'], 1)} a {n(v['deuda_ebitda12'], 1)} veces la utilidad operativa entre 2008 y 2012, lo que "
               "facilita renovar el saldo de 2012. Pero ACC debe comprometerse a cubrir el crédito adicional de hasta "
               f"US${n(v['linea12'])} millones. Si el banco calcula el límite solo sobre el negocio celular, conviene pedir "
               f"un préstamo menor (US${n(v['deuda50'])} millones) y que ACC aporte la diferencia "
               f"(US${n(v['reduccion'])} millones) con recursos propios.", False)])
    numerada([("Antes de firmar. ", True),
              (f"Revisar a fondo las inversiones de ATC en otras empresas, que representan US${n(v['afil'])} millones "
               f"({p(v['afil_sh'], 0)} del valor). Si el precio supera US${n(v['maximo'])} millones, retirarse y proponer a ATC "
               "un acuerdo para alquilarle la fibra de ACC (alternativa D).", False)])

    titulo_seccion("Referencias")
    parrafo([("Stafford, E. y Heilprin, J. (2011, rev. 2012). ", False), ("Valuation of AirThread Connections", False),
             (". Harvard Business School, caso 4263. Cifras: modelo financiero del Grupo 3 (Caso_2_Grupo_3_1.xlsx).", False)],
            tam=9)
    doc.paragraphs[-1].runs[1].italic = True

    salida = Path(salida)
    salida.mkdir(parents=True, exist_ok=True)
    doc.save(salida / "Informe_JD_AirThread.docx")

    # Tabla de control
    lineas = ["# Tabla de control: cifra del informe → celda del Excel del Grupo 3", "",
              "Archivo: `Caso_2_Grupo_3_1.xlsx`. Las cifras derivadas se calculan solo con celdas del archivo.", "",
              "| Clave | Celda | Valor |", "|---|---|---:|"]
    for k, (h, c) in CELDAS.items():
        lineas.append(f"| {k} | `{h}!{c}` | {n(v[k], 4)} |")
    for k, (f, _) in DERIVADAS.items():
        lineas.append(f"| {k} | {f} | {n(v[k], 4)} |")
    (salida / "tabla_control_JD.md").write_text("\n".join(lineas) + "\n", encoding="utf-8")
    print("ok")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
