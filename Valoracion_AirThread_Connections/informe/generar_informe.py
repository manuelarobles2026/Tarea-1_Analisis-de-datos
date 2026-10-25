"""Genera el informe a partir del Excel (única fuente de verdad).

La plantilla usa marcadores [[Hoja!Celda|formato]]. Dos pasos:
  python generar_informe.py control  modelo.xlsx  -> escribe la hoja 'Control' (una fila por cifra)
  (recalcular con LibreOffice)
  python generar_informe.py informe  modelo.xlsx  -> lee los valores de 'Control', arma el .md, el .docx y el .pdf
"""
import os
import re
import subprocess
import sys

from openpyxl import load_workbook
from openpyxl.styles import Font

AQUI = os.path.dirname(os.path.abspath(__file__))
PLANTILLA = os.path.join(AQUI, "plantilla_informe.md")
SALIDA_MD = os.path.join(AQUI, "informe_airthread.md")
SALIDA_CONTROL = os.path.join(AQUI, "tabla_control.md")
MARCA = re.compile(r"\[\[([A-Za-z_]+)!([A-Z]+)(\d+)\|(\w+)\]\]")
FILA0 = 4


def referencias(texto):
    vistas = []
    for m in MARCA.finditer(texto):
        ref = (m.group(1), f"{m.group(2)}{m.group(3)}")
        if ref not in vistas:
            vistas.append(ref)
    return vistas


def descripcion(wb, hoja, celda):
    ws = wb[hoja]
    col = re.match(r"[A-Z]+", celda).group(0)
    fila = int(celda[len(col):])
    ci = ws[celda].column
    etiqueta = None
    for c in range(ci - 1, 0, -1):  # primera etiqueta de texto a la izquierda
        v = ws.cell(row=fila, column=c).value
        if isinstance(v, str) and not v.startswith("="):
            etiqueta = v
            break
    encabezado = None
    for r in range(fila - 1, 0, -1):  # encabezado de columna más cercano hacia arriba
        v = ws.cell(row=r, column=ci).value
        if isinstance(v, str) and not v.startswith("=") and len(v) <= 25:
            encabezado = v
            break
        if isinstance(v, (int, float)) and r < fila - 1 and hoja == "Sensibilidad":
            encabezado = None
            break
    texto = etiqueta or "(sin etiqueta)"
    if encabezado and encabezado != texto:
        texto += f" [{encabezado}]"
    return texto


def paso_control(ruta):
    plantilla = open(PLANTILLA, encoding="utf-8").read()
    refs = referencias(plantilla)
    wb = load_workbook(ruta)
    ws = wb["Control"]
    for r in range(FILA0, ws.max_row + 1):
        for c in "ABCD":
            ws[f"{c}{r}"].value = None
    for i, (hoja, celda) in enumerate(refs):
        r = FILA0 + i
        ws[f"A{r}"] = f"C{i+1:03d}"
        ws[f"B{r}"] = descripcion(wb, hoja, celda)
        ws[f"C{r}"] = f"{hoja}!{celda}"
        ws[f"D{r}"] = f"='{hoja}'!{celda}"
        for c in "ABCD":
            ws[f"{c}{r}"].font = Font(name="Arial", size=10, color="008000" if c == "D" else "000000")
        ws[f"D{r}"].number_format = "#,##0.0000"
    wb.save(ruta)
    print(f"{len(refs)} cifras registradas en la hoja Control")


def num(x, dec):
    s = f"{abs(x):,.{dec}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("−" if x < 0 and round(abs(x), dec) != 0 else "") + s


def formatear(v, fmt):
    if isinstance(v, str):
        return v
    if fmt == "m0":
        return num(v, 0)
    if fmt == "m1":
        return num(v, 1)
    if fmt == "a0":
        return num(abs(v), 0)
    if fmt == "a1":
        return num(abs(v), 1)
    if fmt == "p0":
        return num(v * 100, 0) + "%"
    if fmt == "p1":
        return num(v * 100, 1) + "%"
    if fmt == "p2":
        return num(v * 100, 2) + "%"
    if fmt == "ap1":
        return num(abs(v) * 100, 1) + "%"
    if fmt == "x1":
        return num(v, 1) + "x"
    if fmt == "x2":
        return num(v, 2) + "x"
    if fmt == "pp0":
        return num(v * 100, 0) + " pp"
    if fmt == "m2":
        return num(v, 2)
    if fmt == "n1":
        return num(v, 1)
    if fmt == "n2":
        return num(v, 2)
    if fmt == "n3":
        return num(v, 3)
    if fmt == "n4":
        return num(v, 4)
    if fmt == "sci":
        m, e = f"{v:.1e}".split("e")
        return f"{m.replace('.', ',')} × 10^{int(e)}^"
    if fmt == "ok":
        return v
    raise ValueError(fmt)


def paso_informe(ruta):
    plantilla = open(PLANTILLA, encoding="utf-8").read()
    refs = referencias(plantilla)
    wb = load_workbook(ruta, data_only=True)
    ws = wb["Control"]
    valores = {}
    for i, (hoja, celda) in enumerate(refs):
        r = FILA0 + i
        assert ws[f"C{r}"].value == f"{hoja}!{celda}", (r, ws[f"C{r}"].value, hoja, celda)
        v = ws[f"D{r}"].value
        assert v is not None, f"Control!D{r} vacío: ¿se recalculó?"
        valores[(hoja, celda)] = (v, f"C{i+1:03d}")
    assert wb["Control"]["H18"].value == "OK", "Hay controles de consistencia en REVISAR"

    def sustituir(m):
        v, _ = valores[(m.group(1), f"{m.group(2)}{m.group(3)}")]
        return formatear(v, m.group(4))

    texto = MARCA.sub(sustituir, plantilla)
    open(SALIDA_MD, "w", encoding="utf-8").write(texto)

    # tabla de control
    lineas = ["# Tabla de control: cifra del informe → celda del Excel", "",
              "Cada cifra del informe sale de una celda de `modelo/modelo_airthread.xlsx`. "
              "La misma lista está en la hoja **Control** del Excel (columna D vinculada a la celda).", "",
              "| Clave | Descripción en el Excel | Celda | Valor en el Excel |", "|---|---|---|---:|"]
    wbf = load_workbook(ruta)
    for i, (hoja, celda) in enumerate(refs):
        r = FILA0 + i
        v = ws[f"D{r}"].value
        vv = num(v, 4) if isinstance(v, (int, float)) else str(v)
        desc = wbf["Control"][f"B{r}"].value.replace("|", "/")
        lineas.append(f"| C{i+1:03d} | {desc} | `{hoja}!{celda}` | {vv} |")
    open(SALIDA_CONTROL, "w", encoding="utf-8").write("\n".join(lineas) + "\n")

    # Word (para editar) y PDF (HTML + Chromium, con estilos propios)
    docx = SALIDA_MD.replace(".md", ".docx")
    subprocess.run(["pandoc", SALIDA_MD, "-o", docx, "--from", "markdown+pipe_tables+superscript-tex_math_dollars"], check=True)
    subprocess.run([sys.executable, os.path.join(AQUI, "formato_word.py"), docx], check=True)
    html = SALIDA_MD.replace(".md", ".html")
    subprocess.run(["pandoc", SALIDA_MD, "-s", "-o", html, "--from", "markdown+pipe_tables+superscript-tex_math_dollars",
                    "--css", os.path.join(AQUI, "estilo.css"), "--embed-resources", "--metadata", "lang=es"], check=True)
    pdf = SALIDA_MD.replace(".md", ".pdf")
    subprocess.run(["node", os.path.join(AQUI, "html_a_pdf.js"), html, pdf], check=True)
    os.remove(html)
    print("informe generado:", SALIDA_MD, docx)


if __name__ == "__main__":
    modo, ruta = sys.argv[1], sys.argv[2]
    paso_control(ruta) if modo == "control" else paso_informe(ruta)
