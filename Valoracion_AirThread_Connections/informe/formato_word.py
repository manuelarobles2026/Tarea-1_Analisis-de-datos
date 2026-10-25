"""Da formato al .docx que genera pandoc: Arial, márgenes, tablas con bordes y encabezado azul."""
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = "1F3864"


def borde_celda(celda, color="C9D3E6"):
    tcPr = celda._tc.get_or_add_tcPr()
    bordes = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), color)
        bordes.append(b)
    tcPr.append(bordes)


def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def main(ruta):
    doc = Document(ruta)
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
    for st in doc.styles:
        try:
            st.font.name = "Arial"
            rpr = st.element.get_or_add_rPr()
            rfonts = rpr.find(qn("w:rFonts"))
            if rfonts is None:
                rfonts = OxmlElement("w:rFonts")
                rpr.append(rfonts)
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rfonts.set(qn(a), "Arial")
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
                if rfonts.get(qn(a)) is not None:
                    del rfonts.attrib[qn(a)]
        except (AttributeError, ValueError):
            pass
    estilos = {s.name.lower(): s for s in doc.styles if s.name}
    for nombre, tam in [("normal", 10), ("body text", 10), ("first paragraph", 10), ("compact", 10)]:
        if nombre in estilos:
            estilos[nombre].font.size = Pt(tam)
    for nombre, tam in [("title", 20), ("subtitle", 13), ("heading 1", 14), ("heading 2", 12)]:
        if nombre in estilos:
            st = estilos[nombre]
            st.font.size = Pt(tam)
            st.font.color.rgb = RGBColor.from_string(AZUL)
            st.font.bold = True
    for t in doc.tables:
        t.autofit = True
        for i, fila in enumerate(t.rows):
            for celda in fila.cells:
                borde_celda(celda)
                if i == 0:
                    sombrear(celda, AZUL)
                elif i % 2 == 0:
                    sombrear(celda, "F4F6FB")
                for p in celda.paragraphs:
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    for r in p.runs:
                        r.font.size = Pt(8.5)
                        r.font.name = "Arial"
                        if i == 0:
                            r.font.bold = True
                            r.font.color.rgb = RGBColor(255, 255, 255)
    for p in doc.paragraphs:
        if p.style.name in ("Author", "Date"):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.save(ruta)


if __name__ == "__main__":
    main(sys.argv[1])
