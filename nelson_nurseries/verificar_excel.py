"""Compara los valores recalculados del Excel contra calculos.py.

Uso (después de recalcular el libro con LibreOffice):
    python verificar_excel.py
"""
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as L

from calculos import CAJA_2019, CRECIMIENTOS, ESCENARIOS, flujos, historico, indicadores, proyeccion
from construir_excel import FILA_IND, FILA_VAL, SALIDA

TOL = 1e-6
wb = load_workbook(SALIDA, data_only=True)
r = historico()
k = indicadores(r)
errores = 0
revisadas = 0


def comparar(hoja, ref, esperado, etiqueta):
    global errores, revisadas
    v = wb[hoja][ref].value
    revisadas += 1
    if esperado is None:
        return
    if not isinstance(v, (int, float)) or abs(v - esperado) > TOL * max(1, abs(esperado)):
        errores += 1
        print(f"DIFERENCIA {hoja}!{ref} {etiqueta}: excel={v} python={esperado}")


# Histórico e indicadores
for i, c in enumerate("BCDE"):
    for fila, clave in [(10, "gp"), (13, "ebit"), (15, "ni"), (22, "ca"), (24, "ta"), (29, "cl"), (30, "eq")]:
        comparar("Historico", f"{c}{fila}", r[clave][i], clave)
    for fila, clave in [(5, "crec"), (6, "mb"), (8, "mo"), (9, "tasa"), (10, "mn"), (11, "roa"),
                        (12, "roe"), (13, "dso"), (14, "dio"), (15, "dpo"), (16, "ccc"), (17, "rotaf"),
                        (18, "rotat"), (19, "cash_pct"), (20, "liq"), (21, "acida"), (22, "ktno"),
                        (23, "ktno_pct")]:
        if not (clave == "crec" and i == 0):
            comparar("Indicadores", f"{c}{fila}", k[clave][i], clave)
for f, c in zip(flujos(r), "CDE"):
    for fila, clave in [(11, "cfo"), (13, "fcf"), (14, "d_cash")]:
        comparar("Flujo de caja", f"{c}{fila}", f[clave], clave)

# Proyecciones
for hoja, columnas in [("Proyeccion 2020", ESCENARIOS),
                       ("Sensibilidad", [(None, {"g": g}) for g in CRECIMIENTOS] +
                        [(None, {"g": g, "caja_obj": CAJA_2019}) for g in CRECIMIENTOS])]:
    for j, (_, kw) in enumerate(columnas):
        col = L(2 + j)
        p = proyeccion(r, **kw)
        for clave, fila in {**FILA_VAL, **FILA_IND}.items():
            if clave == "cobertura" and p[clave] is None:
                continue
            comparar(hoja, f"{col}{fila}", p[clave], f"{clave} ({kw})")
        if abs(wb[hoja][f"{col}50"].value) > TOL:
            errores += 1
            print(f"BALANCE NO CUADRA {hoja}!{col}50")

print(f"celdas comparadas: {revisadas}, diferencias: {errores}")
