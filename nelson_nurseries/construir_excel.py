"""Construye el modelo Excel de Nelson Nurseries con fórmulas vivas.

Uso:
    python construir_excel.py            # escribe nelson_nurseries_modelo.xlsx
Luego se recalcula con LibreOffice y verificar_excel.py compara contra calculos.py.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from calculos import BENCH, CAJA_2019, CRECIMIENTOS, ESCENARIOS, EXH2, H

SALIDA = Path(__file__).with_name("nelson_nurseries_modelo.xlsx")

F = "Arial"
AZUL = Font(name=F, color="0000FF")
VERDE = Font(name=F, color="008000")
NEGRO = Font(name=F)
NEGRITA = Font(name=F, bold=True)
TITULO = Font(name=F, bold=True, size=13)
NOTA = Font(name=F, italic=True, size=9, color="555555")
AMARILLO = PatternFill("solid", fgColor="FFFF00")
GRIS = PatternFill("solid", fgColor="E7E6E6")
LINEA = Border(top=Side(style="thin"))

MILES = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
DIAS = '0.0'
VECES = '0.00"x"'

wb = Workbook()


def celda(ws, ref, valor, fuente=NEGRO, fmt=None, relleno=None, borde=None):
    c = ws[ref]
    c.value = valor
    c.font = fuente
    if fmt:
        c.number_format = fmt
    if relleno:
        c.fill = relleno
    if borde:
        c.border = borde
    return c


def encabezado(ws, titulo, subtitulo):
    celda(ws, "A1", titulo, TITULO)
    celda(ws, "A2", subtitulo, NOTA)


# --------------------------------------------------------------------------
# Supuestos
# --------------------------------------------------------------------------
sup = wb.active
sup.title = "Supuestos"
encabezado(sup, "Nelson Nurseries - Supuestos del modelo (miles de USD)",
           "Azul = dato ingresado; negro = fórmula; verde = vínculo a otra hoja; "
           "amarillo = supuesto clave que se puede cambiar.")
SUP = [
    # fila, etiqueta, valor, formato, fuente
    (4, "Crecimiento de ventas 2020", 0.30, PCT,
     "Caso p.2: Roger espera un crecimiento récord de 30% en 2020."),
    (5, "Capex 2020", 75.0, MILES,
     "Caso p.2: Christine espera capex de USD 75 mil (incluye la compra de 12 acres)."),
    (6, "Depreciación 2020", 46.0, MILES, "Caso p.2: depreciación esperada 2020 USD 46 mil."),
    (7, "Tasa de interés de la deuda", 0.065, PCT,
     "Caso p.2: tasa hipotecaria vigente 6,5%. Se usa como costo de la deuda nueva."),
    (8, "Caja objetivo (% de ventas)", 0.08, PCT, "Caso p.2: meta operativa de Christine = 8% de las ventas."),
    (9, "Caja mínima alternativa (% ventas, nivel 2019)", None, PCT,
     "Caja 2019 / ventas 2019 (0,9%): escenario sin reponer la caja."),
    (10, "Tasa efectiva de impuestos (2019)", None, PCT, "Impuestos 2019 / utilidad operativa 2019."),
    (11, "Salario anual de los dueños (en SG&A)", 50.0, MILES,
     "Caso p.2: los Barton retiran USD 50 mil al año como salario, registrado en SG&A."),
    (12, "Descuento por pronto pago proveedores", 0.02, PCT, "Caso p.2: 2% si se paga dentro de 10 días."),
    (13, "Días para obtener el descuento", 10, '0', "Caso p.2: términos 2/10 neto 30."),
    (14, "Plazo neto proveedores (días)", 30, '0', "Caso p.2: términos 2/10 neto 30."),
]
for fila, etiqueta, valor, fmt, fuente_txt in SUP:
    celda(sup, f"A{fila}", etiqueta)
    c = celda(sup, f"B{fila}", valor, AZUL, fmt, AMARILLO if fila in (4, 7, 8) else None)
    celda(sup, f"C{fila}", fuente_txt, NOTA)
sup["B9"].value = "=Historico!E18/Historico!E8"
sup["B9"].font = VERDE
sup["B10"].value = "=Historico!E14/Historico!E13"
sup["B10"].font = VERDE
celda(sup, "A16", "Los escenarios de la hoja 'Proyeccion 2020' toman estos valores salvo cuando "
      "la celda del escenario está en azul con fondo amarillo (supuesto propio del escenario).", NOTA)
sup.column_dimensions["A"].width = 46
sup.column_dimensions["B"].width = 12
sup.column_dimensions["C"].width = 90

# --------------------------------------------------------------------------
# Histórico (misma disposición de filas que el Exhibit 1 del caso)
# --------------------------------------------------------------------------
his = wb.create_sheet("Historico")
encabezado(his, "Estados financieros históricos (miles de USD)",
           "Fuente: Nelson Nurseries (UVA-F-1944), Exhibit 1. Subtotales recalculados con fórmulas.")
COLS = ["B", "C", "D", "E"]
for c, a in zip(COLS, [2016, 2017, 2018, 2019]):
    celda(his, f"{c}6", str(a), NEGRITA).alignment = Alignment(horizontal="right")
FILAS_H = [
    (7, "Estado de resultados", None, True),
    (8, "Ventas", "rev", False), (9, "Costo de ventas", "cogs", False),
    (10, "Utilidad bruta", "=@-@", False), (11, "Gastos de venta y administración (SG&A)", "sga", False),
    (12, "Depreciación", "dep", False), (13, "Utilidad operativa (EBIT)", "=@-@-@", False),
    (14, "Impuestos", "tax", False), (15, "Utilidad neta", "=@-@", False),
    (17, "Balance general", None, True),
    (18, "Caja", "cash", False), (19, "Cuentas por cobrar", "ar", False),
    (20, "Inventario", "inv", False), (21, "Otros activos corrientes", "oca", False),
    (22, "Activo corriente", "SUM", False), (23, "Activo fijo neto", "nfa", False),
    (24, "Activo total", "TA", False),
    (26, "Cuentas por pagar (proveedores)", "ap", False), (27, "Salarios por pagar", "wp", False),
    (28, "Otras cuentas por pagar", "op", False), (29, "Pasivo corriente", "SUMCL", False),
    (30, "Patrimonio (net worth)", "EQ", False),
    (32, "Inversión en activo fijo (capex)", "capex", False), (33, "Compras a proveedores", "purch", False),
]
for fila, etiqueta, clave, es_titulo in FILAS_H:
    celda(his, f"A{fila}", etiqueta, NEGRITA if es_titulo else NEGRO)
    if clave is None:
        continue
    for i, c in enumerate(COLS):
        if clave in H:
            celda(his, f"{c}{fila}", H[clave][i], AZUL, MILES)
            continue
        f = {
            10: f"={c}8-{c}9", 13: f"={c}10-{c}11-{c}12", 15: f"={c}13-{c}14",
            22: f"=SUM({c}18:{c}21)", 24: f"={c}22+{c}23", 29: f"=SUM({c}26:{c}28)",
            30: f"={c}24-{c}29",
        }[fila]
        celda(his, f"{c}{fila}", f, NEGRITA if fila in (15, 24, 30) else NEGRO, MILES,
              borde=LINEA if fila in (10, 13, 15, 22, 24, 29, 30) else None)
celda(his, "A35", "Verificaciones (deben ser 0)", NEGRITA)
celda(his, "A36", "Patrimonio - (patrimonio año anterior + utilidad neta)")
celda(his, "A37", "Activo fijo - (AF anterior + capex - depreciación)")
for i, c in enumerate(COLS[1:], start=1):
    p = COLS[i - 1]
    celda(his, f"{c}36", f"=ROUND({c}30-({p}30+{c}15),6)", NEGRO, MILES)
    celda(his, f"{c}37", f"=ROUND({c}23-({p}23+{c}32-{c}12),6)", NEGRO, MILES)
celda(his, "A38", "Ambas cuadran: no hubo dividendos ni aportes, y el activo fijo no tuvo ventas ni ajustes.", NOTA)
his.column_dimensions["A"].width = 52
for c in COLS:
    his.column_dimensions[c].width = 11

# --------------------------------------------------------------------------
# Indicadores históricos
# --------------------------------------------------------------------------
ind = wb.create_sheet("Indicadores")
encabezado(ind, "Indicadores históricos vs. benchmark",
           "Calculados desde 'Historico'. Columnas H-K: valores publicados en el Exhibit 2; "
           "L-O: diferencia (solo redondeo).")
for c, a in zip(COLS, [2016, 2017, 2018, 2019]):
    celda(ind, f"{c}4", str(a), NEGRITA)
celda(ind, "F4", "Benchmark", NEGRITA)
celda(ind, "H3", "Publicado Exhibit 2", NEGRITA)
celda(ind, "L3", "Diferencia calculado - publicado", NEGRITA)
for j, a in enumerate([2016, 2017, 2018, 2019]):
    celda(ind, f"{L(8 + j)}4", str(a), NEGRITA)
    celda(ind, f"{L(12 + j)}4", str(a), NEGRITA)

h = "Historico!"
IND = [
    # fila, etiqueta, clave, fórmula con {c} y {p} (año previo), formato
    (5, "Crecimiento de ventas", "crec", "={h}{c}8/{h}{p}8-1", PCT),
    (6, "Margen bruto", "mb", "={h}{c}10/{h}{c}8", PCT),
    (7, "SG&A / ventas", None, "={h}{c}11/{h}{c}8", PCT),
    (8, "Margen operativo", "mo", "={h}{c}13/{h}{c}8", PCT),
    (9, "Tasa efectiva de impuestos", None, "={h}{c}14/{h}{c}13", PCT),
    (10, "Margen neto", "mn", "={h}{c}15/{h}{c}8", PCT),
    (11, "ROA (utilidad neta / activo total)", "roa", "={h}{c}15/{h}{c}24", PCT),
    (12, "ROE (utilidad neta / patrimonio)", "roe", "={h}{c}15/{h}{c}30", PCT),
    (13, "Días de cartera (DSO)", "dso", "={h}{c}19/{h}{c}8*365", DIAS),
    (14, "Días de inventario (DIO)", "dio", "={h}{c}20/{h}{c}9*365", DIAS),
    (15, "Días de proveedores (DPO, sobre compras)", "dpo", "={h}{c}26/{h}{c}33*365", DIAS),
    (16, "Ciclo de conversión de caja (DSO+DIO-DPO)", None, "={c}13+{c}14-{c}15", DIAS),
    (17, "Rotación de activo fijo (ventas / AF neto)", "rotaf", "={h}{c}8/{h}{c}23", VECES),
    (18, "Rotación de activo total (ventas / activo total)", None, "={h}{c}8/{h}{c}24", VECES),
    (19, "Caja / ventas (meta 8%)", None, "={h}{c}18/{h}{c}8", PCT),
    (20, "Razón corriente", None, "={h}{c}22/{h}{c}29", VECES),
    (21, "Prueba ácida ((caja + CxC) / pasivo corriente)", None, "=({h}{c}18+{h}{c}19)/{h}{c}29", VECES),
    (22, "Capital de trabajo operativo (CxC+Inv+OAC-PC)", None, "={h}{c}19+{h}{c}20+{h}{c}21-{h}{c}29", MILES),
    (23, "Capital de trabajo operativo / ventas", None, "={c}22/{h}{c}8", PCT),
]
for fila, etiqueta, clave, plantilla, fmt in IND:
    celda(ind, f"A{fila}", etiqueta)
    for i, c in enumerate(COLS):
        if clave == "crec" and i == 0:
            celda(ind, f"{c}{fila}", EXH2["crec"][0], AZUL, fmt).comment = Comment(
                "Dato del Exhibit 2: no se tiene la venta 2015 para recalcularlo.", "modelo")
            continue
        p = COLS[i - 1] if i else c
        celda(ind, f"{c}{fila}", plantilla.format(h=h, c=c, p=p), NEGRO, fmt)
    if clave:
        celda(ind, f"F{fila}", BENCH[clave], AZUL, fmt)
        for j in range(4):
            celda(ind, f"{L(8 + j)}{fila}", EXH2[clave][j], AZUL, fmt)
            celda(ind, f"{L(12 + j)}{fila}", f"={L(2 + j)}{fila}-{L(8 + j)}{fila}", NEGRO,
                  '0.0000' if fmt != PCT else '0.00%')
celda(ind, "F16", "=F13+F14-F15", NEGRO, DIAS)

celda(ind, "A26", "Crecimiento autosostenible 2020 (sin deuda ni aportes, ratios 2019)", NEGRITA)
celda(ind, "A27", "Activo total / ventas 2019")
celda(ind, "E27", "=Historico!E24/Historico!E8", VERDE, '0.000')
celda(ind, "A28", "Pasivo espontáneo / ventas 2019")
celda(ind, "E28", "=Historico!E29/Historico!E8", VERDE, '0.000')
celda(ind, "A29", "Margen neto 2019")
celda(ind, "E29", "=E10", NEGRO, PCT)
celda(ind, "A30", "Crecimiento autosostenible = m / (A/V - P/V - m)")
celda(ind, "E30", "=E29/(E27-E28-E29)", NEGRITA, PCT)
celda(ind, "A31", "Comprobación: ROE sobre patrimonio inicial 2019 (sin dividendos)")
celda(ind, "E31", "=Historico!E15/Historico!D30", NEGRO, PCT)
ind.column_dimensions["A"].width = 52
for j in range(2, 16):
    ind.column_dimensions[L(j)].width = 10

# --------------------------------------------------------------------------
# Flujo de caja histórico
# --------------------------------------------------------------------------
fc = wb.create_sheet("Flujo de caja")
encabezado(fc, "Flujo de caja histórico, método indirecto (miles de USD)",
           "Explica por qué la caja cayó de 120,1 a 9,4 pese a utilidades crecientes.")
for c, a in zip(["C", "D", "E"], [2017, 2018, 2019]):
    celda(fc, f"{c}4", str(a), NEGRITA)
celda(fc, "F4", "Total 17-19", NEGRITA)
FC = [
    (5, "Utilidad neta", "={h}{c}15"),
    (6, "+ Depreciación", "={h}{c}12"),
    (7, "- Aumento cuentas por cobrar", "=-({h}{c}19-{h}{p}19)"),
    (8, "- Aumento inventario", "=-({h}{c}20-{h}{p}20)"),
    (9, "- Aumento otros activos corrientes", "=-({h}{c}21-{h}{p}21)"),
    (10, "+ Aumento pasivo corriente", "={h}{c}29-{h}{p}29"),
    (11, "Flujo de caja operativo", "=SUM({c}5:{c}10)"),
    (12, "- Capex", "=-{h}{c}32"),
    (13, "Flujo de caja libre", "={c}11+{c}12"),
    (14, "Variación de caja según balance", "={h}{c}18-{h}{p}18"),
    (15, "Diferencia (debe ser 0)", "=ROUND({c}13-{c}14,6)"),
]
for fila, etiqueta, plantilla in FC:
    celda(fc, f"A{fila}", etiqueta, NEGRITA if fila in (11, 13) else NEGRO)
    for c, p in zip(["C", "D", "E"], ["B", "C", "D"]):
        celda(fc, f"{c}{fila}", plantilla.format(h=h, c=c, p=p),
              VERDE if fila not in (11, 13, 15) else NEGRITA, MILES,
              borde=LINEA if fila in (11, 13) else None)
    celda(fc, f"F{fila}", f"=SUM(C{fila}:E{fila})", NEGRO, MILES)
celda(fc, "A17", "Uso del flujo operativo antes de cambios en capital de trabajo (utilidad + depreciación)", NEGRITA)
celda(fc, "A18", "Utilidad + depreciación (2017-2019)")
celda(fc, "F18", "=F5+F6", NEGRO, MILES)
celda(fc, "A19", "Absorbido por inventario")
celda(fc, "F19", "=-F8", NEGRO, MILES)
celda(fc, "A20", "Absorbido por cuentas por cobrar")
celda(fc, "F20", "=-F7", NEGRO, MILES)
celda(fc, "A21", "Inventario + cartera como % de utilidad + depreciación")
celda(fc, "F21", "=(F19+F20)/F18", NEGRITA, PCT)
fc.column_dimensions["A"].width = 60
for c in "CDEF":
    fc.column_dimensions[c].width = 12

# --------------------------------------------------------------------------
# Motor de proyección (se usa en dos hojas)
# --------------------------------------------------------------------------
S = "Supuestos!"
ETIQUETAS = {
    6: "Crecimiento de ventas", 7: "Crecimiento para el que se arma inventario y compras",
    8: "Días de cartera (DSO)", 9: "Días de inventario (DIO, sobre costo de ventas)",
    10: "Días de proveedores (DPO, sobre compras)", 11: "Cambio en costo de ventas / ventas (pp)",
    12: "Parte fija del SG&A", 13: "Caja objetivo (% ventas)", 14: "Capex", 15: "Depreciación",
    16: "Tasa de interés deuda nueva", 17: "Tasa de impuestos",
    19: "Estado de resultados 2020", 20: "Ventas", 21: "Costo de ventas", 22: "Utilidad bruta",
    23: "SG&A", 24: "Depreciación", 25: "Utilidad operativa (EBIT)", 26: "Intereses",
    27: "Utilidad antes de impuestos", 28: "Impuestos", 29: "Utilidad neta",
    31: "Balance 2020", 32: "Caja objetivo", 33: "Excedente de caja", 34: "Caja total",
    35: "Cuentas por cobrar", 36: "Inventario", 37: "Otros activos corrientes",
    38: "Activo corriente", 39: "Activo fijo neto", 40: "Activo total",
    42: "Compras a proveedores", 43: "Cuentas por pagar", 44: "Salarios por pagar",
    45: "Otras cuentas por pagar", 46: "Pasivo corriente", 47: "Deuda bancaria requerida",
    48: "Patrimonio", 49: "Pasivo + patrimonio", 50: "Cuadre (debe ser 0)",
    52: "Necesidad de financiamiento externo", 53: "Activos requeridos (con caja objetivo)",
    54: "(-) Pasivo espontáneo", 55: "(-) Patrimonio 2019 + utilidad antes de intereses",
    56: "Necesidad antes de intereses", 57: "Necesidad incluyendo intereses (negativo = excedente)",
    59: "Flujo de caja 2020", 60: "Utilidad neta", 61: "+ Depreciación",
    62: "- Aumento capital de trabajo operativo", 63: "Flujo de caja operativo", 64: "- Capex",
    65: "Flujo de caja libre",
    67: "Indicadores 2020", 68: "Crecimiento de ventas", 69: "Margen bruto", 70: "Margen operativo",
    71: "Margen neto", 72: "ROA", 73: "ROE", 74: "DSO", 75: "DIO", 76: "DPO",
    77: "Ciclo de conversión de caja", 78: "Rotación activo fijo", 79: "Caja / ventas",
    80: "Razón corriente (incluye deuda)", 81: "Deuda / patrimonio", 82: "Cobertura de intereses (EBIT/int.)",
    83: "Deuda / EBITDA",
}
TITULOS = {19, 31, 52, 59, 67}
FILA_IND = dict(crec=68, mb=69, mo=70, mn=71, roa=72, roe=73, dso=74, dio=75, dpo=76, ccc=77,
                rotaf=78, cash_pct=79, liq=80, de=81, cobertura=82, deuda_ebitda=83)
FILA_VAL = dict(rev=20, cogs=21, sga=23, ebit=25, int=26, tax=28, ni=29, cash=34, ar=35, inv=36,
                ta=40, cl=46, debt=47, eq=48, check=50, afn0=56, cfo=63, fcf=65, excedente=33)


def motor(ws, col, nombre, kw):
    """Escribe una columna de proyección. kw = supuestos propios del escenario."""
    c = col
    celda(ws, f"{c}4", nombre, NEGRITA).alignment = Alignment(wrap_text=True, vertical="top")
    # Supuestos: valor por defecto (vínculo) o propio del escenario (azul, amarillo)
    defecto = {
        6: (f"={S}$B$4", PCT), 7: (f"={c}6", PCT),
        8: ("=Historico!$E$19/Historico!$E$8*365", DIAS),
        9: ("=Historico!$E$20/Historico!$E$9*365", DIAS),
        10: ("=Historico!$E$26/Historico!$E$33*365", DIAS),
        11: (0, PCT), 12: (0, MILES), 13: (f"={S}$B$8", PCT), 14: (f"={S}$B$5", MILES),
        15: (f"={S}$B$6", MILES), 16: (f"={S}$B$7", PCT), 17: (f"={S}$B$10", PCT),
    }
    propio = {"g": 6, "g_inv": 7, "dso": 8, "dio": 9, "dpo": 10, "d_cogs": 11, "sga_fijo": 12,
              "caja_obj": 13}
    filas_propias = {propio[k]: v for k, v in kw.items()}
    for fila, (f, fmt) in defecto.items():
        if fila in filas_propias:
            v = filas_propias[fila]
            if fila == 12 and v == 50:
                v = f"={S}$B$11"
            if fila == 13 and abs(v - CAJA_2019) < 1e-12:
                v = f"={S}$B$9"
            celda(ws, f"{c}{fila}", v, AZUL if not str(v).startswith("=") else VERDE, fmt, AMARILLO)
        else:
            fuente = AZUL if not isinstance(f, str) else (VERDE if "!" in f else NEGRO)
            celda(ws, f"{c}{fila}", f, fuente, fmt)
    he = "Historico!$E$"
    fx = {
        20: f"={he}8*(1+{c}6)",
        21: f"={c}20*({he}9/{he}8+{c}11)",
        22: f"={c}20-{c}21",
        23: f"={c}12+{c}20*({he}11-{c}12)/{he}8",
        24: f"={c}15",
        25: f"={c}22-{c}23-{c}24",
        26: f"={c}47*{c}16",
        27: f"={c}25-{c}26",
        28: f"={c}27*{c}17",
        29: f"={c}27-{c}28",
        32: f"={c}13*{c}20",
        33: f"=MAX(-{c}57,0)",
        34: f"={c}32+{c}33",
        35: f"={c}20*{c}8/365",
        36: f"={he}8*(1+{c}7)*({he}9/{he}8+{c}11)*{c}9/365",
        37: f"={c}20*{he}21/{he}8",
        38: f"=SUM({c}34:{c}37)",
        39: f"={he}23+{c}14-{c}15",
        40: f"={c}38+{c}39",
        42: f"={he}33*(1+{c}7)",
        43: f"={c}42*{c}10/365",
        44: f"={c}20*{he}27/{he}8",
        45: f"={c}20*{he}28/{he}8",
        46: f"=SUM({c}43:{c}45)",
        47: f"=MAX({c}57,0)",
        48: f"={he}30+{c}29",
        49: f"={c}46+{c}47+{c}48",
        50: f"=ROUND({c}40-{c}49,6)",
        53: f"={c}32+{c}35+{c}36+{c}37+{c}39",
        54: f"={c}46",
        55: f"={he}30+{c}25*(1-{c}17)",
        56: f"={c}53-{c}54-{c}55",
        57: f"=IF({c}56>0,{c}56/(1-{c}16*(1-{c}17)),{c}56)",
        60: f"={c}29",
        61: f"={c}24",
        62: f"=-(({c}35-{he}19)+({c}36-{he}20)+({c}37-{he}21)-({c}46-{he}29))",
        63: f"={c}60+{c}61+{c}62",
        64: f"=-{c}14",
        65: f"={c}63+{c}64",
        68: f"={c}20/{he}8-1",
        69: f"={c}22/{c}20",
        70: f"={c}25/{c}20",
        71: f"={c}29/{c}20",
        72: f"={c}29/{c}40",
        73: f"={c}29/{c}48",
        74: f"={c}35/{c}20*365",
        75: f"={c}36/{c}21*365",
        76: f"={c}43/{c}42*365",
        77: f"={c}74+{c}75-{c}76",
        78: f"={c}20/{c}39",
        79: f"={c}34/{c}20",
        80: f"={c}38/({c}46+{c}47)",
        81: f"={c}47/{c}48",
        82: f'=IF({c}26>0,{c}25/{c}26,"sin deuda")',
        83: f"={c}47/({c}25+{c}24)",
    }
    fmts = {68: PCT, 69: PCT, 70: PCT, 71: PCT, 72: PCT, 73: PCT, 74: DIAS, 75: DIAS, 76: DIAS,
            77: DIAS, 78: VECES, 79: PCT, 80: VECES, 81: VECES, 82: VECES, 83: VECES}
    fuertes = {22, 25, 29, 38, 40, 46, 49, 50, 57, 63, 65, 47}
    for fila, f in fx.items():
        cc = celda(ws, f"{c}{fila}", f, NEGRITA if fila in fuertes else NEGRO, fmts.get(fila, MILES),
                   borde=LINEA if fila in (22, 25, 29, 38, 40, 46, 49, 56, 63, 65) else None)
        if fila in (47, 57):
            cc.fill = GRIS


def hoja_proyeccion(nombre, titulo, nota, columnas):
    ws = wb.create_sheet(nombre)
    encabezado(ws, titulo, nota)
    celda(ws, "A3", "Supuestos de cada escenario (azul con fondo amarillo = distinto de la base)", NOTA)
    celda(ws, "A5", "Supuestos", NEGRITA)
    for fila, etq in ETIQUETAS.items():
        celda(ws, f"A{fila}", etq, NEGRITA if fila in TITULOS else NEGRO)
    for j, (nom, kw) in enumerate(columnas):
        motor(ws, L(2 + j), nom, kw)
        ws.column_dimensions[L(2 + j)].width = 15
    ws.column_dimensions["A"].width = 52
    ws.row_dimensions[4].height = 58
    ws.freeze_panes = "B5"
    return ws


hoja_proyeccion(
    "Proyeccion 2020", "Proyección 2020 por escenario (miles de USD)",
    "Base = se mantienen los ratios de 2019 y se crece 30%. La deuda paga 6,5% sobre el saldo de cierre; "
    "se resuelve en forma cerrada (fila 57) para evitar referencias circulares.",
    ESCENARIOS)

COLS_SENS = ([(f"Crec. {g:.0%}\ncaja 8%", {"g": g}) for g in CRECIMIENTOS] +
             [(f"Crec. {g:.0%}\ncaja 0,9%", {"g": g, "caja_obj": CAJA_2019}) for g in CRECIMIENTOS])
hoja_proyeccion(
    "Sensibilidad", "Sensibilidad de la necesidad de financiamiento al crecimiento (miles de USD)",
    "Ratios de 2019. Columnas B-I: reponiendo la caja a 8% de ventas; J-Q: manteniendo la caja en 0,9% de ventas.",
    COLS_SENS)

# --------------------------------------------------------------------------
# Descuento de proveedores
# --------------------------------------------------------------------------
dp = wb.create_sheet("Descuento proveedores")
encabezado(dp, "¿Conviene dejar de tomar el descuento 2/10 neto 30?",
           "Compara el costo implícito de financiarse con proveedores contra la tasa bancaria.")
DP = [
    (4, "Descuento", f"={S}B12", PCT, VERDE),
    (5, "Días para obtener descuento", f"={S}B13", '0', VERDE),
    (6, "Plazo neto", f"={S}B14", '0', VERDE),
    (7, "Costo anual simple de no tomar el descuento", "=B4/(1-B4)*365/(B6-B5)", PCT, NEGRITA),
    (8, "Costo anual efectivo (compuesto)", "=(1+B4/(1-B4))^(365/(B6-B5))-1", PCT, NEGRITA),
    (9, "Tasa de la deuda bancaria / hipotecaria", f"={S}B7", PCT, VERDE),
    (10, "Compras 2020 (escenario base)", "='Proyeccion 2020'!B42", MILES, VERDE),
    (11, "Ahorro anual por tomar el descuento", "=B10*B4", MILES, NEGRO),
    (12, "Financiamiento adicional si se paga al día 30", "=B10*(B6-B5)/365", MILES, NEGRO),
    (13, "Costo de esa misma suma con deuda bancaria", "=B12*B9", MILES, NEGRO),
]
for fila, etq, f, fmt, fuente in DP:
    celda(dp, f"A{fila}", etq)
    celda(dp, f"B{fila}", f, fuente, fmt)
celda(dp, "A15", "Conclusión: seguir tomando el descuento (pagar el día 10, no antes). Dejar de tomarlo "
      "liberaría poca caja a un costo implícito de ~37% anual, muy superior a la deuda bancaria.", NOTA)
dp.column_dimensions["A"].width = 52
dp.column_dimensions["B"].width = 12

# --------------------------------------------------------------------------
# Resumen
# --------------------------------------------------------------------------
res = wb.create_sheet("Resumen", 0)
encabezado(res, "Nelson Nurseries - Resumen de resultados (miles de USD)",
           "Todas las cifras están vinculadas a las hojas de cálculo; cambiar 'Supuestos' actualiza todo.")
celda(res, "A4", "Escenario", NEGRITA)
for j, t in enumerate(["Ventas", "Utilidad neta", "Margen neto", "ROE", "Flujo caja libre",
                       "Deuda requerida", "Deuda/patrimonio", "Cobertura int."]):
    celda(res, f"{L(2 + j)}4", t, NEGRITA).alignment = Alignment(wrap_text=True)
filas_res = [(20, MILES), (29, MILES), (71, PCT), (73, PCT), (65, MILES), (47, MILES), (81, VECES),
             (82, VECES)]
for i, (nom, _) in enumerate(ESCENARIOS):
    col = L(2 + i)
    celda(res, f"A{5 + i}", f"='Proyeccion 2020'!{col}4", VERDE)
    for j, (fila, fmt) in enumerate(filas_res):
        celda(res, f"{L(2 + j)}{5 + i}", f"='Proyeccion 2020'!{col}{fila}", VERDE, fmt)
celda(res, "A15", "Referencias", NEGRITA)
celda(res, "A16", "Crecimiento autosostenible (sin financiamiento externo)")
celda(res, "B16", "=Indicadores!E30", VERDE, PCT)
celda(res, "A17", "Flujo de caja libre acumulado 2017-2019")
celda(res, "B17", "='Flujo de caja'!F13", VERDE, MILES)
celda(res, "A18", "Utilidad neta acumulada 2017-2019")
celda(res, "B18", "='Flujo de caja'!F5", VERDE, MILES)
celda(res, "A19", "Costo anual de no tomar el descuento de proveedores")
celda(res, "B19", "='Descuento proveedores'!B7", VERDE, PCT)
res.column_dimensions["A"].width = 46
for j in range(2, 10):
    res.column_dimensions[L(j)].width = 13
res.row_dimensions[4].height = 30

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False

if __name__ == "__main__":
    wb.save(SALIDA)
    print("escrito", SALIDA)
