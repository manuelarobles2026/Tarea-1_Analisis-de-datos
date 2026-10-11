"""Construye el modelo en Excel del caso Cafés Monte Bianco (HBS 204-S18).

Todas las cifras calculadas quedan como fórmulas de Excel; los datos del caso
son las únicas cifras escritas a mano (en azul).

Uso: python construir_excel.py <ruta_salida.xlsx>
"""
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL, NEGRO, VERDE = "0000FF", "000000", "008000"
FUENTE = "Arial"
F_MILES = '#,##0;(#,##0);"-"'
F_MILES1 = '#,##0.0;(#,##0.0);"-"'
F_PRECIO = '#,##0.00;(#,##0.00);"-"'
F_PCT = '0.0%;(0.0%);"-"'
F_KG = '#,##0;(#,##0);"-"'
F_X = '0.00"x";(0.00"x");"-"'
F_DIAS = '0.0;(0.0);"-"'

RELLENO_TITULO = PatternFill("solid", fgColor="1F3864")
RELLENO_SECCION = PatternFill("solid", fgColor="D9E1F2")
RELLENO_CLAVE = PatternFill("solid", fgColor="FFFF00")
RELLENO_OPT = PatternFill("solid", fgColor="E2EFDA")
LINEA = Side(style="thin", color="808080")

MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio",
         "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
CALIDADES = ["D", "C", "B", "BB", "A", "AA", "AAA"]
PRECIOS = [7.00, 15.50, 21.00, 24.00, 28.00, 31.00, 34.00]
COSTO_UNIT = [5.25, 9.93, 11.36, 12.96, 14.15, 15.25, 16.26]
COSTO_FIJO = [2641, 3262, 3262, 3262, 3262, 3262, 3262]
VOLUMENES = [  # Anexo 2: filas = % marketing 0..10 %, columnas = C..AAA
    [2395000, 1645000, 1529000, 1103000, 667000, 127000],
    [2455000, 1696000, 1566000, 1145000, 717000, 157000],
    [2545000, 1720700, 1649000, 1226000, 799000, 229000],
    [2665000, 1896000, 1760000, 1334000, 910000, 325000],
    [2800000, 2024000, 1884000, 1456000, 1034000, 433000],
    [2950000, 2166000, 2022000, 1591000, 1172000, 553000],
    [3130000, 2367000, 2157000, 1723000, 1307000, 667000],
    [3196000, 2417000, 2195000, 1792000, 1345000, 694000],
    [3268000, 2446000, 2231000, 1823000, 1366000, 752000],
    [3310000, 2465000, 2258000, 1837000, 1495000, 875000],
    [3340000, 2494000, 2286000, 1843000, 1523000, 935000],
]
ESTACIONALIDAD = [0.071, 0.090, 0.118, 0.075, 0.081, 0.050, 0.042, 0.033,
                  0.068, 0.131, 0.121, 0.120]

class _Faltantes(dict):
    """Referencias hacia adelante: se rellenan con 0 y la celda se reescribe después."""

    def __missing__(self, clave):
        return "0"


N = {}  # nombre lógico -> referencia absoluta 'Hoja'!$C$5


def ref(hoja, celda):
    col = "".join(ch for ch in celda if ch.isalpha())
    fila = "".join(ch for ch in celda if ch.isdigit())
    return f"{hoja}!${col}${fila}"


def put(ws, celda, valor, fmt=None, negrita=False, nota=None, relleno=None,
        cursiva=False, alinear=None):
    c = ws[celda]
    c.value = valor
    if isinstance(valor, str) and valor.startswith("="):
        color = VERDE if "!" in valor else NEGRO
    elif isinstance(valor, (int, float)):
        color = AZUL
    else:
        color = NEGRO
    c.font = Font(name=FUENTE, size=10, bold=negrita, color=color, italic=cursiva)
    if fmt:
        c.number_format = fmt
    if nota:
        c.comment = Comment(nota, "Analista")
    if relleno:
        c.fill = relleno
    if alinear:
        c.alignment = Alignment(horizontal=alinear, vertical="center", wrap_text=True)
    return c


def titulo(ws, texto, subtitulo=None, ancho=8):
    ws["A1"].value = texto
    ws["A1"].font = Font(name=FUENTE, size=13, bold=True, color="FFFFFF")
    for col in range(1, ancho + 1):
        ws.cell(row=1, column=col).fill = RELLENO_TITULO
    if subtitulo:
        ws["A2"].value = subtitulo
        ws["A2"].font = Font(name=FUENTE, size=9, italic=True, color="595959")


def seccion(ws, fila, texto, ancho=8):
    ws.cell(row=fila, column=1).value = texto
    ws.cell(row=fila, column=1).font = Font(name=FUENTE, size=10, bold=True)
    for col in range(1, ancho + 1):
        ws.cell(row=fila, column=col).fill = RELLENO_SECCION


def encabezados(ws, fila, textos, col_ini=1):
    for i, t in enumerate(textos):
        c = ws.cell(row=fila, column=col_ini + i)
        c.value = t
        c.font = Font(name=FUENTE, size=10, bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=LINEA)


def total(ws, fila, cols):
    for col in cols:
        c = ws[f"{col}{fila}"]
        c.border = Border(top=LINEA)
        c.font = Font(name=FUENTE, size=10, bold=True, color=c.font.color, italic=c.font.italic)


def anchos(ws, d):
    for col, w in d.items():
        ws.column_dimensions[col].width = w


# --------------------------------------------------------------------------
# Hoja Supuestos
# --------------------------------------------------------------------------
def hoja_supuestos(wb):
    ws = wb.active
    ws.title = "Supuestos"
    titulo(ws, "Cafés Monte Bianco – Supuestos y datos del caso (miles de €, salvo indicación)",
           "Azul = dato del caso o supuesto editable · Negro = fórmula · Verde = vínculo a otra hoja · "
           "Amarillo = supuesto clave", 4)
    encabezados(ws, 3, ["Concepto", "Valor", "Unidad", "Fuente / comentario"])
    filas = [
        ("SEC", "1. Estado de resultados 2019 (Anexo 1)"),
        ("v_mp19", "Ingresos marca propia 2019", 7902.72, F_MILES1, "miles €", "Anexo 1"),
        ("v_pr19", "Ingresos marca de calidad superior 2019", 36705.24, F_MILES1, "miles €", "Anexo 1"),
        ("cmv19", "Costo de mercadería vendida 2019", 26421, F_MILES, "miles €", "Anexo 1 (incluye depreciación)"),
        ("dep19", "Depreciación 2019 incluida en el CMV", 2081, F_MILES, "miles €", "Nota (*) del Anexo 1"),
        ("mkt19", "Gastos de marketing 2019", 3303, F_MILES, "miles €", "Anexo 1"),
        ("id19", "Gastos de I+D 2019", 2647, F_MILES, "miles €", "Anexo 1"),
        ("vta19", "Gastos de venta 2019", 2844, F_MILES, "miles €", "Anexo 1"),
        ("adm19", "Gastos administrativos 2019", 3781, F_MILES, "miles €", "Anexo 1"),
        ("int19", "Gastos de intereses 2019", 3043, F_MILES, "miles €", "Anexo 1"),
        ("tasa_imp", "Tasa de impuestos", 0.40, F_PCT, "%", "Anexo 1"),
        ("SEC", "2. Balance al 31-dic-2019 (Anexo 1)"),
        ("caja19", "Caja", 892, F_MILES, "miles €", "Anexo 1"),
        ("invmp19", "Inventario de materia prima", 2313, F_MILES, "miles €", "Anexo 1"),
        ("invpt19", "Bienes terminados", 913, F_MILES, "miles €", "Anexo 1"),
        ("invpt_kg19", "Bienes terminados (kg)", 174000, F_KG, "kg", "Anexo 1 y Anexo 5 (inventario inicial)"),
        ("cxc19", "Cuentas a cobrar", 7448.928, F_MILES1, "miles €", "Anexo 1"),
        ("ppe19", "Bienes de uso (bruto, incluye la ampliación de 4,5 M€)", 33706, F_MILES, "miles €", "Anexo 1"),
        ("depac19", "Depreciación acumulada", 9758, F_MILES, "miles €", "Anexo 1 (se resta)"),
        ("lote", "Lote desocupado incluido en bienes de uso", 10000, F_MILES, "miles €",
         "Nota (*) del balance; se asume en miles de € como el resto del estado"),
        ("cxp19", "Cuentas a pagar", 388, F_MILES, "miles €", "Anexo 1"),
        ("lc19", "Línea de crédito", 20000, F_MILES, "miles €", "Anexo 1"),
        ("dlp19", "Deuda a largo plazo", 8000, F_MILES, "miles €", "Anexo 1"),
        ("lim_lc", "Límite total de la línea de crédito", 20000, F_MILES, "miles €", "Texto del caso (Dino)"),
        ("SEC", "3. Volúmenes 2019 (Anexo 4)"),
        ("kg_mp19", "Kilos vendidos marca propia 2019", 1152000, F_KG, "kg", "Anexo 4"),
        ("kg_pr19", "Kilos vendidos calidad superior 2019", 1196000, F_KG, "kg", "Anexo 4"),
        ("SEC", "4. Capacidad e inversión"),
        ("cap_mes", "Capacidad de producción mensual tras la ampliación", 500000, F_KG, "kg/mes",
         "350.000 + 150.000 kg/mes (texto del caso y Anexo 5)"),
        ("cap", "Capacidad anual", "=B{cap_mes}*12", F_KG, "kg/año", "6.000.000 kg (Anexo 5)"),
        ("inv_amp", "Inversión de la ampliación (ya incluida en bienes de uso 2019)", 4500, F_MILES, "miles €", "Texto del caso"),
        ("vida", "Vida útil de la ampliación", 15, "0", "años", "Texto del caso"),
        ("dep_inc", "Depreciación incremental anual", "=B{inv_amp}/B{vida}", F_MILES, "miles €", "Lineal"),
        ("dep_tot", "Depreciación total 2020", "=B{dep19}+B{dep_inc}", F_MILES, "miles €",
         "2.081 base + incremental (indicación del enunciado)"),
        ("cf_incl_dep", "¿Los costos fijos del Anexo 3 ya incluyen la depreciación base? (1 = sí, 0 = no)", 1, "0", "1/0",
         "Supuesto: sí, porque la depreciación 2019 está dentro del CMV; con 0 se suma 2.081 aparte al CMV"),
        ("SEC", "5. Marca propia (Anexos 2 y 3)"),
        ("p_mp", "Precio marca propia", "=Anexos!B4", F_PRECIO, "€/kg", "Anexo 2 / texto (7,00 €/kg)"),
        ("cv_mp", "Costo variable unitario marca propia", "=Anexos!B5", F_PRECIO, "€/kg", "Anexo 3"),
        ("cf_mp", "Costos fijos de manufactura solo marca propia", "=Anexos!B6", F_MILES, "miles €", "Anexo 3"),
        ("cf_pr", "Costos fijos de manufactura con calidad superior (incluye los de marca propia)", "=Anexos!C6",
         F_MILES, "miles €", "Anexo 3 y su nota (*)"),
        ("SEC", "6. Ahorros de gastos si se pasa 100 % a marca propia (Carla)"),
        ("ah_mkt", "Ahorro en marketing", 1.00, F_PCT, "%", "Texto del caso"),
        ("ah_vta", "Ahorro en gastos de venta", 0.65, F_PCT, "%", "Texto del caso"),
        ("ah_id", "Ahorro en I+D", 0.75, F_PCT, "%", "Texto del caso"),
        ("ah_adm", "Ahorro en gastos administrativos", 0.50, F_PCT, "%", "Texto del caso"),
        ("SEC", "7. Ahorros de gastos en el escenario de mix (marca propia + calidad superior)"),
        ("ahx_vta", "Ahorro en gastos de venta (mix)", 0.0, F_PCT, "%",
         "Supuesto: se mantiene la estructura comercial de 2019 porque sigue el negocio de calidad superior"),
        ("ahx_id", "Ahorro en I+D (mix)", 0.0, F_PCT, "%", "Supuesto: se mantiene I+D (calidad superior lo requiere)"),
        ("ahx_adm", "Ahorro en gastos administrativos (mix)", 0.0, F_PCT, "%", "Supuesto: se mantiene"),
        ("SEC", "8. Capital de trabajo y financiamiento"),
        ("meses_cxc_mp", "Meses de ventas en CxC al cierre – marca propia (pago a 90 días)", 3, "0", "meses",
         "CxC final = ventas de oct + nov + dic (Anexo 4 y enunciado)"),
        ("meses_cxc_pr", "Meses de ventas en CxC al cierre – calidad superior (pago a 30 días)", 1, "0", "meses",
         "CxC final = ventas de diciembre (política de 30 días)"),
        ("sh_cxc_mp", "Proporción de ventas anuales en CxC – marca propia",
         "=SUM(OFFSET(Anexos!$B$36,1-B{meses_cxc_mp},0,B{meses_cxc_mp},1))", F_PCT, "%",
         "Suma de los últimos meses según la estacionalidad del Anexo 4"),
        ("sh_cxc_pr", "Proporción de ventas anuales en CxC – calidad superior",
         "=SUM(OFFSET(Anexos!$B$36,1-B{meses_cxc_pr},0,B{meses_cxc_pr},1))", F_PCT, "%",
         "Suma de los últimos meses según la estacionalidad del Anexo 4"),
        ("rot_invmp", "Rotación inventario de materia prima 2019 (CMV sin depreciación / inventario)",
         "=(B{cmv19}-B{dep19})/B{invmp19}", F_X, "veces", "Se mantiene en 2020"),
        ("rot_cxp", "Rotación de cuentas a pagar 2019 (CMV sin depreciación / CxP)",
         "=(B{cmv19}-B{dep19})/B{cxp19}", F_X, "veces", "Se mantiene en 2020"),
        ("caja_vtas", "Caja mínima operativa (% de ventas, igual que 2019)", "=B{caja19}/(B{v_mp19}+B{v_pr19})",
         F_PCT, "%", "La caja del balance proyectado es la variable de cierre; se compara contra este mínimo"),
        ("invpt_kg20", "Inventario final de bienes terminados 2020", 0, F_KG, "kg", "Anexo 5 (inventario final cero)"),
        ("tasa_int", "Tasa de interés implícita 2019", "=B{int19}/(B{lc19}+B{dlp19})", F_PCT, "%",
         "Intereses / deuda total; la deuda se mantiene constante en 2020"),
        ("divid", "Dividendos 2020", 0, F_MILES, "miles €", "Supuesto: no se reparten (toda la utilidad se retiene)"),
    ]
    fila = 4
    pendientes = []
    for item in filas:
        if item[0] == "SEC":
            fila += 1
            seccion(ws, fila, item[1], 4)
            fila += 1
            continue
        clave, etiqueta, valor, fmt, unidad, fuente = item
        N[clave] = ref("Supuestos", f"B{fila}")
        put(ws, f"A{fila}", etiqueta)
        pendientes.append((fila, valor, fmt))
        put(ws, f"C{fila}", unidad)
        put(ws, f"D{fila}", fuente, cursiva=True)
        ws[f"D{fila}"].font = Font(name=FUENTE, size=9, italic=True, color="595959")
        fila += 1
    filas_por_clave = {k: v.split("$")[-1] for k, v in N.items()}
    for f, valor, fmt in pendientes:
        if isinstance(valor, str):
            valor = valor.format(**filas_por_clave)
        put(ws, f"B{f}", valor, fmt)
    for clave in ("cf_incl_dep", "meses_cxc_mp", "meses_cxc_pr", "ahx_vta", "ahx_id", "ahx_adm"):
        ws[N[clave].split("!")[1].replace("$", "")].fill = RELLENO_CLAVE
    anchos(ws, {"A": 74, "B": 14, "C": 10, "D": 88})
    ws.freeze_panes = "A4"


# --------------------------------------------------------------------------
# Hoja Anexos
# --------------------------------------------------------------------------
def hoja_anexos(wb):
    ws = wb.create_sheet("Anexos")
    titulo(ws, "Datos de los Anexos 2, 3 y 4", "Fuente: caso HBS 204-S18, Cafés Monte Bianco", 8)
    encabezados(ws, 3, ["Calidad del café"] + CALIDADES)
    put(ws, "A4", "Precio por kilo (€)")
    put(ws, "A5", "Costo unitario variable (€/kg)")
    put(ws, "A6", "Costos fijos (miles €) *")
    for i in range(7):
        col = get_column_letter(2 + i)
        put(ws, f"{col}4", PRECIOS[i], F_PRECIO)
        put(ws, f"{col}5", COSTO_UNIT[i], F_PRECIO)
        put(ws, f"{col}6", COSTO_FIJO[i], F_MILES)
    put(ws, "A7", "* Los costos fijos de calidad superior incluyen los costos fijos de marca propia (Anexo 3).",
        cursiva=True)
    seccion(ws, 8, "Anexo 2 – Volumen anual estimado (kg) según % de gasto de marketing / ventas")
    encabezados(ws, 9, ["Marketing / ventas"] + CALIDADES)
    for i in range(11):
        f = 10 + i
        put(ws, f"A{f}", i / 100, F_PCT)
        if i == 0:
            put(ws, f"B{f}", 6000000, F_KG, nota="Marca propia: 6.000.000 kg a 7 €/kg sin gasto de marketing")
        for j in range(6):
            put(ws, f"{get_column_letter(3 + j)}{f}", VOLUMENES[i][j], F_KG)
    seccion(ws, 23, "Anexo 4 – Ventas 2019 por mes")
    encabezados(ws, 24, ["Mes", "% del total anual", "Marca propia (kg)", "Calidad superior (kg)"])
    for i, m in enumerate(MESES):
        f = 25 + i
        put(ws, f"A{f}", m)
        put(ws, f"B{f}", ESTACIONALIDAD[i], F_PCT)
        put(ws, f"C{f}", f"=Supuestos!$B${N['kg_mp19'].split('$')[-1]}*B{f}", F_KG)
        put(ws, f"D{f}", f"=Supuestos!$B${N['kg_pr19'].split('$')[-1]}*B{f}", F_KG)
    put(ws, "A37", "Total", negrita=True)
    put(ws, "B37", "=SUM(B25:B36)", F_PCT)
    put(ws, "C37", "=SUM(C25:C36)", F_KG)
    put(ws, "D37", "=SUM(D25:D36)", F_KG)
    total(ws, 37, "BCD")
    anchos(ws, {"A": 34, **{get_column_letter(c): 14 for c in range(2, 9)}})


# --------------------------------------------------------------------------
# Hoja Diagnóstico 2019
# --------------------------------------------------------------------------
def hoja_diagnostico(wb):
    ws = wb.create_sheet("Diagnostico_2019")
    S = lambda k: N[k]  # noqa: E731
    titulo(ws, "Diagnóstico financiero 2019 – ¿fue un año “excelente”?", "Miles de €", 4)
    seccion(ws, 3, "Estado de resultados 2019", 4)
    encabezados(ws, 4, ["Concepto", "Miles €", "% de ventas"])
    er = [
        ("Ingresos marca propia", f"={S('v_mp19')}"),
        ("Ingresos calidad superior", f"={S('v_pr19')}"),
        ("Ingresos totales", "=B5+B6"),
        ("Costo de mercadería vendida (incluye depreciación)", f"={S('cmv19')}"),
        ("Margen bruto", "=B7-B8"),
        ("Gastos de marketing", f"={S('mkt19')}"),
        ("Gastos de I+D", f"={S('id19')}"),
        ("Gastos de venta", f"={S('vta19')}"),
        ("Gastos administrativos", f"={S('adm19')}"),
        ("Utilidad operativa (EBIT)", "=B9-SUM(B10:B13)"),
        ("Gastos de intereses", f"={S('int19')}"),
        ("Utilidad antes de impuestos", "=B14-B15"),
        ("Impuestos", f"=MAX(0,B16)*{S('tasa_imp')}"),
        ("Utilidad neta", "=B16-B17"),
        ("Depreciación (dentro del CMV)", f"={S('dep19')}"),
        ("EBITDA", "=B14+B19"),
    ]
    for i, (et, fo) in enumerate(er):
        f = 5 + i
        put(ws, f"A{f}", et, negrita=et in ("Ingresos totales", "Margen bruto", "Utilidad operativa (EBIT)", "Utilidad neta"))
        put(ws, f"B{f}", fo, F_MILES)
        put(ws, f"C{f}", f"=B{f}/$B$7", F_PCT)
    for f in (7, 9, 14, 18):
        total(ws, f, "ABC")

    seccion(ws, 22, "Balance al 31-dic-2019", 4)
    encabezados(ws, 23, ["Concepto", "Miles €", "% del activo"])
    bal = [
        ("Caja", f"={S('caja19')}"),
        ("Inventario de materia prima", f"={S('invmp19')}"),
        ("Bienes terminados", f"={S('invpt19')}"),
        ("Cuentas a cobrar", f"={S('cxc19')}"),
        ("Activo corriente", "=SUM(B24:B27)"),
        ("Bienes de uso (bruto)", f"={S('ppe19')}"),
        ("Depreciación acumulada", f"=-{S('depac19')}"),
        ("Bienes de uso netos", "=B29+B30"),
        ("Total activos", "=B28+B31"),
        ("Cuentas a pagar", f"={S('cxp19')}"),
        ("Línea de crédito", f"={S('lc19')}"),
        ("Pasivo corriente", "=B33+B34"),
        ("Deuda a largo plazo", f"={S('dlp19')}"),
        ("Total pasivos", "=B35+B36"),
        ("Patrimonio neto", "=B32-B37"),
        ("Total pasivos y patrimonio", "=B37+B38"),
    ]
    for i, (et, fo) in enumerate(bal):
        f = 24 + i
        put(ws, f"A{f}", et, negrita=et.startswith("Total") or et in ("Activo corriente", "Pasivo corriente", "Patrimonio neto"))
        put(ws, f"B{f}", fo, F_MILES)
        put(ws, f"C{f}", f"=B{f}/$B$32", F_PCT)
    for f in (28, 32, 35, 37, 39):
        total(ws, f, "ABC")
    put(ws, "A40", "Nota: el total del activo suma 35.515 (el PDF del caso muestra 33.515 por error de tipeo; "
        "el patrimonio de 7.127 cuadra con 35.515).", cursiva=True)

    seccion(ws, 42, "Indicadores clave 2019", 4)
    encabezados(ws, 43, ["Indicador", "Valor", "Lectura"])
    ind = [
        ("Rentabilidad", None, None, None),
        ("Margen bruto", "=B9/B7", F_PCT, "Holgado por el peso de calidad superior"),
        ("Margen operativo (EBIT / ventas)", "=B14/B7", F_PCT, ""),
        ("Margen neto", "=B18/B7", F_PCT, "Muy bajo: los intereses se comen más de la mitad del EBIT"),
        ("Intereses / EBIT", "=B15/B14", F_PCT, ""),
        ("Cobertura de intereses (EBIT / intereses)", "=B14/B15", F_X, "Menos de 2x: zona de riesgo"),
        ("Rotación de activos (ventas / activos)", "=B7/B32", F_X, ""),
        ("Multiplicador de capital (activos / patrimonio)", "=B32/B38", F_X, "Apalancamiento muy alto"),
        ("ROA (utilidad neta / activos)", "=B18/B32", F_PCT, ""),
        ("ROE (utilidad neta / patrimonio)", "=B18/B38", F_PCT, "Alto, pero explicado por el apalancamiento"),
        ("ROE por DuPont (margen × rotación × multiplicador)", "=B47*B50*B51", F_PCT, "Comprobación"),
        ("Descomposición del ROE por apalancamiento", None, None, None),
        ("RONA después de impuestos = EBIT(1–t) / (deuda + patrimonio)",
         f"=B14*(1-{S('tasa_imp')})/(B34+B36+B38)", F_PCT, "Rentabilidad operativa sin efecto de la deuda"),
        ("Costo de la deuda después de impuestos", f"={S('tasa_int')}*(1-{S('tasa_imp')})", F_PCT,
         f"Tasa implícita antes de impuestos ≈ intereses / deuda"),
        ("Deuda financiera / patrimonio", "=(B34+B36)/B38", F_X, ""),
        ("ROE = RONA + (RONA – kd(1–t)) × D/E", "=B56+(B56-B57)*B58", F_PCT,
         "Casi 12 de los 22 puntos del ROE vienen de la deuda"),
        ("ROE si no hubiera deuda (= RONA)", "=B56", F_PCT, ""),
        ("Capital de trabajo y liquidez", None, None, None),
        ("Días de cobro (CxC / ventas × 365)", "=B27/B7*365", F_DIAS, "Promedio; marca propia paga a 90 días"),
        ("Días de inventario MP (sobre CMV sin depreciación)", "=B25/(B8-B19)*365", F_DIAS, ""),
        ("Días de bienes terminados", "=B26/(B8-B19)*365", F_DIAS, ""),
        ("Días de pago (CxP / CMV sin depreciación × 365)", "=B33/(B8-B19)*365", F_DIAS, "Casi no hay crédito de proveedores"),
        ("Ciclo de conversión de efectivo (días)", "=B62+B63+B64-B65", F_DIAS, ""),
        ("Capital de trabajo operativo (inventarios + CxC – CxP)", "=B25+B26+B27-B33", F_MILES, ""),
        ("Capital de trabajo operativo / ventas", "=B67/B7", F_PCT, "Cada € de venta adicional exige ~0,25 € de capital"),
        ("Razón corriente (activo corriente / pasivo corriente)", "=B28/B35", F_X, "Menor que 1"),
        ("Prueba ácida ((caja + CxC) / pasivo corriente)", "=(B24+B27)/B35", F_X, ""),
        ("Línea de crédito disponible", f"={S('lim_lc')}-B34", F_MILES, "Totalmente utilizada"),
        ("Deuda total / EBITDA", "=(B34+B36)/B20", F_X, ""),
        ("Lote desocupado / total activos", f"={S('lote')}/B32", F_PCT, "Activo improductivo financiado con deuda"),
        ("Comprobación de CxC con la estacionalidad (Anexo 4)", None, None, None),
        ("CxC estimada: calidad superior (diciembre)", f"=B6*{S('sh_cxc_pr')}", F_MILES, ""),
        ("CxC estimada: marca propia (oct + nov + dic)", f"=B5*{S('sh_cxc_mp')}", F_MILES, ""),
        ("CxC estimada total vs. CxC del balance", "=B75+B76", F_MILES, "Coincide razonablemente con 7.449"),
        ("Precios y márgenes unitarios 2019", None, None, None),
        ("Precio promedio marca propia (€/kg)", f"=B5/{S('kg_mp19')}*1000", F_PRECIO, "Contra 7,00 €/kg en 2020"),
        ("Precio promedio calidad superior (€/kg)", f"=B6/{S('kg_pr19')}*1000", F_PRECIO, "Entre las calidades A y AA"),
        ("Ventas totales 2019 (kg)", f"={S('kg_mp19')}+{S('kg_pr19')}", F_KG, ""),
        ("Uso de la capacidad 2019 (350.000 kg/mes)", "=B81/(350000*12)", F_PCT, "Capacidad ociosa antes de la ampliación"),
        ("Uso de la capacidad 2020 con el mismo volumen", f"=B81/{S('cap')}", F_PCT, ""),
    ]
    f = 44
    for et, fo, fmt, lec in ind:
        if fo is None:
            put(ws, f"A{f}", et, negrita=True, cursiva=True)
        else:
            put(ws, f"A{f}", et)
            put(ws, f"B{f}", fo, fmt)
            put(ws, f"C{f}", lec, cursiva=True)
        f += 1
    anchos(ws, {"A": 62, "B": 14, "C": 62})
    ws.freeze_panes = "A3"


# --------------------------------------------------------------------------
# Hoja Mix de producción
# --------------------------------------------------------------------------
def hoja_mix(wb):
    ws = wb.create_sheet("Mix_Produccion")
    S = lambda k: N[k]  # noqa: E731
    titulo(ws, "Mini estado de resultados por alternativa de mix: (p – cv)·q – CF – %M·p·q",
           "Miles de €. Capacidad = 6.000.000 kg; la capacidad no usada por calidad superior se llena con marca propia "
           "a 7 €/kg (demanda garantizada a plena capacidad).", 9)
    put(ws, "A3", "Capacidad anual (kg)")
    put(ws, "B3", f"={S('cap')}", F_KG)
    put(ws, "A4", "Precio marca propia (€/kg)")
    put(ws, "B4", f"={S('p_mp')}", F_PRECIO)
    put(ws, "A5", "Costo variable marca propia (€/kg)")
    put(ws, "B5", f"={S('cv_mp')}", F_PRECIO)
    put(ws, "A6", "Costos fijos solo marca propia")
    put(ws, "B6", f"={S('cf_mp')}", F_MILES)
    put(ws, "A7", "Costos fijos con calidad superior (incluyen los de marca propia)")
    put(ws, "B7", f"={S('cf_pr')}", F_MILES)
    put(ws, "A8", "Resultado solo marca propia (D) a plena capacidad")
    put(ws, "B8", "=(B4-B5)*B3/1000-B6", F_MILES, negrita=True)

    cols = [get_column_letter(3 + j) for j in range(6)]  # C..H  = C..AAA, igual que en Anexos

    def cabecera(fila):
        put(ws, f"A{fila}", "Calidad", negrita=True)
        put(ws, f"A{fila + 1}", "Precio (€/kg)")
        put(ws, f"A{fila + 2}", "Costo variable (€/kg)")
        for c in cols:
            put(ws, f"{c}{fila}", f"=Anexos!{c}3", negrita=True, alinear="center")
            put(ws, f"{c}{fila + 1}", f"=Anexos!{c}4", F_PRECIO)
            put(ws, f"{c}{fila + 2}", f"=Anexos!{c}5", F_PRECIO)

    def tabla(fila_sec, texto, formula, fmt):
        seccion(ws, fila_sec, texto, 9)
        cabecera(fila_sec + 1)
        put(ws, f"A{fila_sec + 4}", "% marketing / ventas ↓", negrita=True)
        for i in range(11):
            f = fila_sec + 5 + i
            put(ws, f"A{f}", f"=Anexos!A{10 + i}", F_PCT)
            for c in cols:
                q = f"Anexos!{c}{10 + i}"
                put(ws, f"{c}{f}", formula(c, f, q, fila_sec + 1), fmt)
        return fila_sec + 5  # primera fila de datos

    # Tabla A: solo calidad superior
    fa = tabla(10, "Tabla A – Resultado vendiendo SOLO calidad superior (miles €)",
               lambda c, f, q, h: f"=({c}${h + 1}-{c}${h + 2}-$A{f}*{c}${h + 1})*{q}/1000-$B$7", F_MILES)
    # Tabla B: capacidad ociosa si solo calidad superior
    fb = tabla(28, "Tabla B – Capacidad ociosa si se vende SOLO calidad superior (% de 6.000.000 kg)",
               lambda c, f, q, h: f"=1-{q}/$B$3", F_PCT)
    # Tabla C: mix con marca propia
    fc = tabla(46, "Tabla C – Resultado del MIX: calidad superior + marca propia en la capacidad restante (miles €)",
               lambda c, f, q, h: (f"=({c}${h + 1}-{c}${h + 2}-$A{f}*{c}${h + 1})*{q}/1000"
                                   f"+($B$4-$B$5)*($B$3-{q})/1000-$B$7"), F_MILES)

    from openpyxl.formatting.rule import FormulaRule
    rango_c = f"C{fc}:H{fc + 10}"
    ws.conditional_formatting.add(rango_c, FormulaRule(
        formula=[f"C{fc}=MAX($C${fc}:$H${fc + 10})"], fill=RELLENO_OPT, font=Font(bold=True)))

    # Resumen por calidad
    r = 64
    seccion(ws, r, "Resumen – mejor combinación por calidad", 9)
    put(ws, f"A{r + 1}", "Calidad", negrita=True)
    etiquetas = [
        ("Mejor resultado del MIX (miles €)", F_MILES),
        ("% marketing óptimo en el MIX", F_PCT),
        ("Kilos de calidad superior", F_KG),
        ("Kilos de marca propia", F_KG),
        ("Mejor resultado SOLO calidad superior (miles €)", F_MILES),
        ("% marketing óptimo SOLO calidad superior", F_PCT),
        ("Kilos SOLO calidad superior", F_KG),
        ("Capacidad ociosa SOLO calidad superior", F_PCT),
        ("Mejora del MIX vs. solo marca propia (miles €)", F_MILES),
    ]
    for i, (et, _) in enumerate(etiquetas):
        put(ws, f"A{r + 2 + i}", et)
    a0, a1 = fa, fa + 10
    c0, c1 = fc, fc + 10
    for idx, c in enumerate(cols):
        anexo_rng = f"Anexos!{c}$10:{c}$20"
        put(ws, f"{c}{r + 1}", f"=Anexos!{c}3", negrita=True, alinear="center")
        put(ws, f"{c}{r + 2}", f"=MAX({c}{c0}:{c}{c1})", F_MILES)
        put(ws, f"{c}{r + 3}", f"=INDEX($A${c0}:$A${c1},MATCH({c}{r + 2},{c}{c0}:{c}{c1},0))", F_PCT)
        put(ws, f"{c}{r + 4}", f"=INDEX({anexo_rng},MATCH({c}{r + 2},{c}{c0}:{c}{c1},0))", F_KG)
        put(ws, f"{c}{r + 5}", f"=$B$3-{c}{r + 4}", F_KG)
        put(ws, f"{c}{r + 6}", f"=MAX({c}{a0}:{c}{a1})", F_MILES)
        put(ws, f"{c}{r + 7}", f"=INDEX($A${a0}:$A${a1},MATCH({c}{r + 6},{c}{a0}:{c}{a1},0))", F_PCT)
        put(ws, f"{c}{r + 8}", f"=INDEX({anexo_rng},MATCH({c}{r + 6},{c}{a0}:{c}{a1},0))", F_KG)
        put(ws, f"{c}{r + 9}", f"=1-{c}{r + 8}/$B$3", F_PCT)
        put(ws, f"{c}{r + 10}", f"={c}{r + 2}-$B$8", F_MILES)

    # Óptimo global
    o = 77
    seccion(ws, o, "MIX ÓPTIMO (máximo de la Tabla C)", 9)
    opt = [
        ("Resultado máximo (miles €)", f"=MAX(C{r + 2}:H{r + 2})", F_MILES),
        ("Calidad de alta gama elegida", f"=INDEX(C{r + 1}:H{r + 1},MATCH(B{o + 1},C{r + 2}:H{r + 2},0))", None),
        ("% de marketing / ventas", f"=INDEX(C{r + 3}:H{r + 3},MATCH(B{o + 1},C{r + 2}:H{r + 2},0))", F_PCT),
        ("Kilos de calidad superior", f"=INDEX(C{r + 4}:H{r + 4},MATCH(B{o + 1},C{r + 2}:H{r + 2},0))", F_KG),
        ("Kilos de marca propia", f"=B3-B{o + 4}", F_KG),
        ("Precio calidad superior (€/kg)", f"=INDEX(C12:H12,MATCH(B{o + 2},C11:H11,0))", F_PRECIO),
        ("Costo variable calidad superior (€/kg)", f"=INDEX(C13:H13,MATCH(B{o + 2},C11:H11,0))", F_PRECIO),
        ("Peso de calidad superior en el volumen", f"=B{o + 4}/B3", F_PCT),
    ]
    for i, (et, fo, fmt) in enumerate(opt):
        put(ws, f"A{o + 1 + i}", et, negrita=(i < 2))
        put(ws, f"B{o + 1 + i}", fo, fmt, negrita=(i < 2), relleno=RELLENO_OPT)
    for k, fila in (("opt_res", o + 1), ("opt_cal", o + 2), ("opt_m", o + 3), ("opt_qpr", o + 4),
                    ("opt_qmp", o + 5), ("opt_p", o + 6), ("opt_cv", o + 7)):
        N[k] = ref("Mix_Produccion", f"B{fila}")

    # Mini estados comparativos
    m = 88
    seccion(ws, m, "Mini estados de resultados comparativos (miles €)", 9)
    encabezados(ws, m + 1, ["Concepto", "Solo marca propia", "Solo AAA (mejor %M)",
                            "Solo calidad superior (mejor)", "MIX óptimo"])
    aaa = "H"
    best_solo = f"MATCH(MAX(C{r + 6}:H{r + 6}),C{r + 6}:H{r + 6},0)"
    escen = {
        # (kg premium, precio, cv, %M, kg marca propia, CF)
        "B": ("0", "0", "0", "0", "$B$3", "$B$6"),
        "C": (f"{aaa}{r + 8}", f"{aaa}12", f"{aaa}13", f"{aaa}{r + 7}", "0", "$B$7"),
        "D": (f"INDEX(C{r + 8}:H{r + 8},{best_solo})", f"INDEX(C12:H12,{best_solo})",
              f"INDEX(C13:H13,{best_solo})", f"INDEX(C{r + 7}:H{r + 7},{best_solo})", "0", "$B$7"),
        "E": (f"B{o + 4}", f"B{o + 6}", f"B{o + 7}", f"B{o + 3}", f"B{o + 5}", "$B$7"),
    }
    filas_mini = ["Calidad superior: kilos", "Calidad superior: precio (€/kg)", "Calidad superior: costo variable (€/kg)",
                  "% marketing / ventas", "Marca propia: kilos", "Ventas calidad superior", "Ventas marca propia",
                  "Ventas totales", "Costos variables", "Margen bruto de contribución (p – cv)·q",
                  "Costos fijos de manufactura", "Publicidad y ventas (%M·p·q)", "Resultado", "Capacidad utilizada"]
    for i, et in enumerate(filas_mini):
        put(ws, f"A{m + 2 + i}", et, negrita=et in ("Ventas totales", "Resultado"))
    for col, (q, p, cv, mm, qmp, cf) in escen.items():
        b = m + 2
        put(ws, f"{col}{b}", f"={q}", F_KG)
        put(ws, f"{col}{b + 1}", f"={p}", F_PRECIO)
        put(ws, f"{col}{b + 2}", f"={cv}", F_PRECIO)
        put(ws, f"{col}{b + 3}", f"={mm}", F_PCT)
        put(ws, f"{col}{b + 4}", f"={qmp}", F_KG)
        put(ws, f"{col}{b + 5}", f"={col}{b}*{col}{b + 1}/1000", F_MILES)
        put(ws, f"{col}{b + 6}", f"={col}{b + 4}*$B$4/1000", F_MILES)
        put(ws, f"{col}{b + 7}", f"={col}{b + 5}+{col}{b + 6}", F_MILES)
        put(ws, f"{col}{b + 8}", f"=({col}{b}*{col}{b + 2}+{col}{b + 4}*$B$5)/1000", F_MILES)
        put(ws, f"{col}{b + 9}", f"={col}{b + 7}-{col}{b + 8}", F_MILES)
        put(ws, f"{col}{b + 10}", f"={cf}", F_MILES)
        put(ws, f"{col}{b + 11}", f"={col}{b + 3}*{col}{b + 5}", F_MILES)
        put(ws, f"{col}{b + 12}", f"={col}{b + 9}-{col}{b + 10}-{col}{b + 11}", F_MILES, negrita=True)
        put(ws, f"{col}{b + 13}", f"=({col}{b}+{col}{b + 4})/$B$3", F_PCT)
    total(ws, m + 9, "ABCDE")
    total(ws, m + 14, "ABCDE")
    # Del resultado del mix a la utilidad neta: gastos de estructura, depreciación incremental e intereses
    extra = ["Depreciación incremental de la ampliación", "I+D, ventas y administración",
             "Intereses (deuda constante)", "Utilidad antes de impuestos", "Utilidad neta (40 % de impuestos)"]
    for i, et in enumerate(extra):
        put(ws, f"A{m + 16 + i}", et, negrita=(i == 4))
    estructura_mp = (f"{S('id19')}*(1-{S('ah_id')})+{S('vta19')}*(1-{S('ah_vta')})+{S('adm19')}*(1-{S('ah_adm')})")
    estructura_pr = (f"{S('id19')}*(1-{S('ahx_id')})+{S('vta19')}*(1-{S('ahx_vta')})+{S('adm19')}*(1-{S('ahx_adm')})")
    for col in "BCDE":
        b = m + 16
        put(ws, f"{col}{b}", f"={S('dep_inc')}", F_MILES)
        put(ws, f"{col}{b + 1}", "=" + (estructura_mp if col == "B" else estructura_pr), F_MILES)
        put(ws, f"{col}{b + 2}", f"={S('tasa_int')}*({S('lc19')}+{S('dlp19')})", F_MILES)
        put(ws, f"{col}{b + 3}", f"={col}{m + 14}-{col}{b}-{col}{b + 1}-{col}{b + 2}", F_MILES)
        put(ws, f"{col}{b + 4}", f"={col}{b + 3}-MAX(0,{col}{b + 3})*{S('tasa_imp')}", F_MILES, negrita=True)
    total(ws, m + 20, "ABCDE")
    put(ws, f"A{m + 22}", "Solo marca propia aplica los ahorros de Carla en I+D, ventas y administración; las demás columnas "
        "mantienen la estructura de 2019 (supuestos de la hoja Supuestos).", cursiva=True)
    put(ws, f"A{m + 23}", "Supuesto simplificador (enunciado): el mix es marca propia + UNA sola calidad de alta gama. "
        "La fila «Resultado» es la fórmula del enunciado; las filas siguientes agregan estructura, depreciación incremental e intereses (ver EEFF_2020).", cursiva=True)

    # Gráfico
    ch = BarChart()
    ch.type = "col"
    ch.title = "Mejor resultado por calidad (miles €)"
    ch.y_axis.title = "Miles €"
    ch.add_data(Reference(ws, min_col=3, max_col=8, min_row=r + 2, max_row=r + 2), from_rows=True, titles_from_data=False)
    ch.add_data(Reference(ws, min_col=3, max_col=8, min_row=r + 6, max_row=r + 6), from_rows=True, titles_from_data=False)
    ch.set_categories(Reference(ws, min_col=3, max_col=8, min_row=r + 1, max_row=r + 1))
    from openpyxl.chart.series import SeriesLabel
    ch.series[0].tx = SeriesLabel(v="MIX con marca propia")
    ch.series[1].tx = SeriesLabel(v="Solo calidad superior")
    ch.series[0].graphicalProperties.solidFill = "2A78D6"
    ch.series[1].graphicalProperties.solidFill = "EB6834"
    ch.height, ch.width = 8, 16
    ws.add_chart(ch, "J64")
    anchos(ws, {"A": 52, "B": 16, **{c: 15 for c in cols}})
    ws.freeze_panes = "B3"


# --------------------------------------------------------------------------
# Hoja EEFF 2020
# --------------------------------------------------------------------------
def hoja_eeff(wb):
    ws = wb.create_sheet("EEFF_2020")
    S = lambda k: N[k]  # noqa: E731
    titulo(ws, "Estados financieros proyectados 2020 – comparación de estrategias (miles €)",
           "Deuda constante (línea de crédito 20.000 + largo plazo 8.000); la caja es la variable de cierre del balance.", 5)
    encabezados(ws, 3, ["Concepto", "2019 Real", "2020 Solo marca propia", "2020 MIX óptimo", "Comentario"])
    R = {}
    fila = [4]

    def linea(clave, etiqueta, b, c, d, fmt=F_MILES, negrita=False, coment="", sec=None):
        if sec:
            fila[0] += 1
            seccion(ws, fila[0], sec, 5)
            fila[0] += 1
        f = fila[0]
        R[clave] = f
        put(ws, f"A{f}", etiqueta, negrita=negrita)
        for col, v in (("B", b), ("C", c), ("D", d)):
            if v is None:
                continue
            if isinstance(v, str):
                v = v.format_map(_Faltantes({**{k: str(x) for k, x in R.items()}, "col": col}))
            put(ws, f"{col}{f}", v, fmt, negrita=negrita)
        if coment:
            put(ws, f"E{f}", coment, cursiva=True)
            ws[f"E{f}"].font = Font(name=FUENTE, size=9, italic=True, color="595959")
        fila[0] += 1
        return f

    # Drivers
    linea("cal", "Calidad de alta gama", None, "—", f"={S('opt_cal')}", None, sec="Supuestos operativos")
    ws[f"B{R['cal']}"].value = "Mix 2019"
    for col in "BCD":
        ws[f"{col}{R['cal']}"].alignment = Alignment(horizontal="right")
    linea("kg_mp", "Kilos marca propia", f"={S('kg_mp19')}", f"={S('cap')}", f"={S('opt_qmp')}", F_KG)
    linea("p_mp", "Precio marca propia (€/kg)", f"={S('v_mp19')}/B{{kg_mp}}*1000", f"={S('p_mp')}", f"={S('p_mp')}", F_PRECIO)
    linea("cv_mp", "Costo variable marca propia (€/kg)", None, f"={S('cv_mp')}", f"={S('cv_mp')}", F_PRECIO)
    linea("kg_pr", "Kilos calidad superior", f"={S('kg_pr19')}", 0, f"={S('opt_qpr')}", F_KG)
    linea("p_pr", "Precio calidad superior (€/kg)", f"={S('v_pr19')}/B{{kg_pr}}*1000", 0, f"={S('opt_p')}", F_PRECIO)
    linea("cv_pr", "Costo variable calidad superior (€/kg)", None, 0, f"={S('opt_cv')}", F_PRECIO)
    linea("pm", "% marketing / ventas de calidad superior", f"={S('mkt19')}/{S('v_pr19')}", 0, f"={S('opt_m')}", F_PCT)
    linea("cf", "Costos fijos de manufactura (Anexo 3)", None, f"={S('cf_mp')}", f"={S('cf_pr')}", F_MILES)
    linea("uso", "Uso de la capacidad de 6.000.000 kg", "=(B{kg_mp}+B{kg_pr})/" + S('cap'),
          "=({col}{kg_mp}+{col}{kg_pr})/" + S('cap'), "=({col}{kg_mp}+{col}{kg_pr})/" + S('cap'), F_PCT)

    # Estado de resultados
    linea("v_mp", "Ventas marca propia", f"={S('v_mp19')}", "={col}{kg_mp}*{col}{p_mp}/1000",
          "={col}{kg_mp}*{col}{p_mp}/1000", sec="Estado de resultados")
    linea("v_pr", "Ventas calidad superior", f"={S('v_pr19')}", "={col}{kg_pr}*{col}{p_pr}/1000",
          "={col}{kg_pr}*{col}{p_pr}/1000")
    linea("vtas", "Ventas totales", "={col}{v_mp}+{col}{v_pr}", "={col}{v_mp}+{col}{v_pr}",
          "={col}{v_mp}+{col}{v_pr}", negrita=True)
    linea("cvmp", "Costo variable marca propia", None, "={col}{kg_mp}*{col}{cv_mp}/1000", "={col}{kg_mp}*{col}{cv_mp}/1000")
    linea("cvpr", "Costo variable calidad superior", None, "={col}{kg_pr}*{col}{cv_pr}/1000", "={col}{kg_pr}*{col}{cv_pr}/1000")
    linea("cfx", "Costos fijos de manufactura (incluyen depreciación base 2.081)", None, "={col}{cf}", "={col}{cf}")
    linea("dinc", "Depreciación incremental de la ampliación (4.500 / 15)", None, f"={S('dep_inc')}", f"={S('dep_inc')}")
    linea("dbase", "Depreciación base adicional (solo si el supuesto = 0)", None,
          f"=(1-{S('cf_incl_dep')})*{S('dep19')}", f"=(1-{S('cf_incl_dep')})*{S('dep19')}")
    linea("cmv", "Costo de mercadería vendida", f"={S('cmv19')}", "=SUM({col}{cvmp}:{col}{dbase})",
          "=SUM({col}{cvmp}:{col}{dbase})")
    linea("mb", "Margen bruto", "={col}{vtas}-{col}{cmv}", "={col}{vtas}-{col}{cmv}", "={col}{vtas}-{col}{cmv}", negrita=True)
    linea("mkt", "Gastos de marketing", f"={S('mkt19')}", f"={S('mkt19')}*(1-{S('ah_mkt')})",
          "={col}{pm}*{col}{v_pr}", coment="Mix: %M × ventas de calidad superior")
    linea("id", "Gastos de I+D", f"={S('id19')}", f"={S('id19')}*(1-{S('ah_id')})", f"={S('id19')}*(1-{S('ahx_id')})")
    linea("vta", "Gastos de venta", f"={S('vta19')}", f"={S('vta19')}*(1-{S('ah_vta')})", f"={S('vta19')}*(1-{S('ahx_vta')})")
    linea("adm", "Gastos administrativos", f"={S('adm19')}", f"={S('adm19')}*(1-{S('ah_adm')})",
          f"={S('adm19')}*(1-{S('ahx_adm')})")
    linea("gop", "Total gastos operativos", "=SUM({col}{mkt}:{col}{adm})", "=SUM({col}{mkt}:{col}{adm})",
          "=SUM({col}{mkt}:{col}{adm})")
    linea("ebit", "Utilidad operativa (EBIT)", "={col}{mb}-{col}{gop}", "={col}{mb}-{col}{gop}",
          "={col}{mb}-{col}{gop}", negrita=True)
    linea("int", "Gastos de intereses", f"={S('int19')}", f"={S('tasa_int')}*({S('lc19')}+{S('dlp19')})",
          f"={S('tasa_int')}*({S('lc19')}+{S('dlp19')})", coment="Deuda constante × tasa implícita 2019")
    linea("uai", "Utilidad antes de impuestos", "={col}{ebit}-{col}{int}", "={col}{ebit}-{col}{int}", "={col}{ebit}-{col}{int}")
    linea("imp", "Impuestos (40 %)", f"=MAX(0,{{col}}{{uai}})*{S('tasa_imp')}", f"=MAX(0,{{col}}{{uai}})*{S('tasa_imp')}",
          f"=MAX(0,{{col}}{{uai}})*{S('tasa_imp')}")
    linea("un", "Utilidad neta", "={col}{uai}-{col}{imp}", "={col}{uai}-{col}{imp}", "={col}{uai}-{col}{imp}", negrita=True)
    linea("dep", "Memo: depreciación total del año", f"={S('dep19')}", f"={S('dep_tot')}", f"={S('dep_tot')}")
    linea("ebitda", "Memo: EBITDA", "={col}{ebit}+{col}{dep}", "={col}{ebit}+{col}{dep}", "={col}{ebit}+{col}{dep}")
    for k in ("vtas", "mb", "ebit", "un"):
        total(ws, R[k], "ABCD")

    # Balance
    linea("caja", "Caja (variable de cierre)", f"={S('caja19')}", "={col}{tpp}-SUM({col}{invmp}:{col}{cxc})-{col}{ppen}",
          "={col}{tpp}-SUM({col}{invmp}:{col}{cxc})-{col}{ppen}", sec="Balance general al 31 de diciembre",
          coment="Negativa = faltante que habría que financiar")
    # placeholders se completan después (referencias hacia adelante)
    caja_row = R["caja"]
    linea("invmp", "Inventario de materia prima", f"={S('invmp19')}",
          f"=({{col}}{{cmv}}-{{col}}{{dep}})/{S('rot_invmp')}", f"=({{col}}{{cmv}}-{{col}}{{dep}})/{S('rot_invmp')}",
          coment="Rotación 2019 sobre CMV sin depreciación")
    linea("invpt", "Bienes terminados", f"={S('invpt19')}", f"={S('invpt_kg20')}*{S('cv_mp')}/1000",
          f"={S('invpt_kg20')}*{S('cv_mp')}/1000", coment="Anexo 5: inventario final cero")
    linea("cxc", "Cuentas a cobrar", f"={S('cxc19')}",
          f"={{col}}{{v_mp}}*{S('sh_cxc_mp')}+{{col}}{{v_pr}}*{S('sh_cxc_pr')}",
          f"={{col}}{{v_mp}}*{S('sh_cxc_mp')}+{{col}}{{v_pr}}*{S('sh_cxc_pr')}",
          coment="Marca propia: oct+nov+dic; calidad superior: dic")
    linea("ac", "Activo corriente", "=SUM({col}{caja}:{col}{cxc})", "=SUM({col}{caja}:{col}{cxc})",
          "=SUM({col}{caja}:{col}{cxc})", negrita=True)
    linea("ppe", "Bienes de uso (bruto)", f"={S('ppe19')}", "=$B{ppe}", "=$B{ppe}".replace("{ppe}", "{ppe}"),
          coment="Sin nuevas inversiones; la ampliación ya está en 2019")
    ws[f"C{R['ppe']}"].value = f"=$B${R['ppe']}"
    ws[f"D{R['ppe']}"].value = f"=$B${R['ppe']}"
    linea("depac", "Depreciación acumulada", f"=-{S('depac19')}", "=$B{depac}-{col}{dep}", "=$B{depac}-{col}{dep}",
          coment="9.758 + 2.081 + 300")
    linea("ppen", "Bienes de uso netos", "={col}{ppe}+{col}{depac}", "={col}{ppe}+{col}{depac}", "={col}{ppe}+{col}{depac}")
    linea("ta", "Total activos", "={col}{ac}+{col}{ppen}", "={col}{ac}+{col}{ppen}", "={col}{ac}+{col}{ppen}", negrita=True)
    linea("cxp", "Cuentas a pagar", f"={S('cxp19')}", f"=({{col}}{{cmv}}-{{col}}{{dep}})/{S('rot_cxp')}",
          f"=({{col}}{{cmv}}-{{col}}{{dep}})/{S('rot_cxp')}")
    linea("lc", "Línea de crédito", f"={S('lc19')}", "=$B{lc}", "=$B{lc}")
    linea("pc", "Pasivo corriente", "={col}{cxp}+{col}{lc}", "={col}{cxp}+{col}{lc}", "={col}{cxp}+{col}{lc}")
    linea("dlp", "Deuda a largo plazo", f"={S('dlp19')}", "=$B{dlp}", "=$B{dlp}")
    linea("tp", "Total pasivos", "={col}{pc}+{col}{dlp}", "={col}{pc}+{col}{dlp}", "={col}{pc}+{col}{dlp}")
    linea("pat", "Patrimonio neto", "={col}{ta}-{col}{tp}", f"=$B{{pat}}+{{col}}{{un}}-{S('divid')}",
          f"=$B{{pat}}+{{col}}{{un}}-{S('divid')}", negrita=True,
          coment="2019: activos – pasivos (7.127); 2020: + utilidad – dividendos")
    linea("tpp", "Total pasivos y patrimonio", "={col}{tp}+{col}{pat}", "={col}{tp}+{col}{pat}", "={col}{tp}+{col}{pat}",
          negrita=True)
    linea("chk", "Control: activos – pasivos y patrimonio", "={col}{ta}-{col}{tpp}", "={col}{ta}-{col}{tpp}",
          "={col}{ta}-{col}{tpp}")
    # rehacer caja con las filas ya conocidas
    for col in "CD":
        ws[f"{col}{caja_row}"].value = (f"={col}{R['tpp']}-SUM({col}{R['invmp']}:{col}{R['cxc']})-{col}{R['ppen']}")
    for k in ("ac", "ta", "tpp"):
        total(ws, R[k], "ABCD")
    linea("cajamin", "Memo: caja mínima operativa (% de ventas 2019)", f"={{col}}{{vtas}}*{S('caja_vtas')}",
          f"={{col}}{{vtas}}*{S('caja_vtas')}", f"={{col}}{{vtas}}*{S('caja_vtas')}")
    linea("nec", "Memo: excedente / (necesidad) de financiamiento", "={col}{caja}-{col}{cajamin}",
          "={col}{caja}-{col}{cajamin}", "={col}{caja}-{col}{cajamin}", negrita=True,
          coment="Negativo = fondos adicionales requeridos con la deuda congelada")
    linea("lcdisp", "Memo: línea de crédito disponible", f"={S('lim_lc')}-{{col}}{{lc}}", f"={S('lim_lc')}-{{col}}{{lc}}",
          f"={S('lim_lc')}-{{col}}{{lc}}")
    linea("intad", "Memo: interés adicional si la necesidad se cubre con deuda (antes de impuestos)", None,
          f"=MAX(0,-{{col}}{{nec}})*{S('tasa_int')}", f"=MAX(0,-{{col}}{{nec}})*{S('tasa_int')}")
    linea("unaj", "Memo: utilidad neta ajustada por ese interés adicional", None,
          f"={{col}}{{un}}-{{col}}{{intad}}*(1-{S('tasa_imp')})", f"={{col}}{{un}}-{{col}}{{intad}}*(1-{S('tasa_imp')})")

    # Flujo de efectivo
    linea("f_un", "Utilidad neta", None, "={col}{un}", "={col}{un}", sec="Estado de flujo de efectivo 2020 (método indirecto)")
    linea("f_dep", "(+) Depreciación", None, "={col}{dep}", "={col}{dep}")
    linea("f_cxc", "(–) Aumento de cuentas a cobrar", None, "=-({col}{cxc}-$B{cxc})", "=-({col}{cxc}-$B{cxc})")
    linea("f_imp", "(–) Aumento de inventario de materia prima", None, "=-({col}{invmp}-$B{invmp})", "=-({col}{invmp}-$B{invmp})")
    linea("f_ipt", "(–) Aumento de bienes terminados", None, "=-({col}{invpt}-$B{invpt})", "=-({col}{invpt}-$B{invpt})")
    linea("f_cxp", "(+) Aumento de cuentas a pagar", None, "={col}{cxp}-$B{cxp}", "={col}{cxp}-$B{cxp}")
    linea("fco", "Flujo de caja operativo", None, "=SUM({col}{f_un}:{col}{f_cxp})", "=SUM({col}{f_un}:{col}{f_cxp})", negrita=True)
    linea("capex", "(–) Inversiones en bienes de uso", None, "=-({col}{ppe}-$B{ppe})", "=-({col}{ppe}-$B{ppe})")
    linea("fci", "Flujo de caja de inversión", None, "={col}{capex}", "={col}{capex}", negrita=True)
    linea("f_deu", "(+/–) Variación de deuda", None, "=({col}{lc}+{col}{dlp})-($B{lc}+$B{dlp})",
          "=({col}{lc}+{col}{dlp})-($B{lc}+$B{dlp})")
    linea("f_div", "(–) Dividendos", None, f"=-{S('divid')}", f"=-{S('divid')}")
    linea("fcf", "Flujo de caja de financiamiento", None, "={col}{f_deu}+{col}{f_div}", "={col}{f_deu}+{col}{f_div}", negrita=True)
    linea("dcaja", "Variación neta de caja", None, "={col}{fco}+{col}{fci}+{col}{fcf}", "={col}{fco}+{col}{fci}+{col}{fcf}",
          negrita=True)
    linea("cini", "Caja inicial", None, "=$B{caja}", "=$B{caja}")
    linea("cfin", "Caja final", None, "={col}{cini}+{col}{dcaja}", "={col}{cini}+{col}{dcaja}", negrita=True)
    linea("chk2", "Control: caja final – caja del balance", None, "={col}{cfin}-{col}{caja}", "={col}{cfin}-{col}{caja}")
    for k in ("fco", "dcaja", "cfin"):
        total(ws, R[k], "ABCD")

    # Indicadores
    linea("r_mb", "Margen bruto", "={col}{mb}/{col}{vtas}", "={col}{mb}/{col}{vtas}", "={col}{mb}/{col}{vtas}", F_PCT,
          sec="Indicadores (rueda de ganancias, rueda de caja y rueda del ROE)")
    linea("r_ebit", "Margen operativo (EBIT / ventas)", "={col}{ebit}/{col}{vtas}", "={col}{ebit}/{col}{vtas}",
          "={col}{ebit}/{col}{vtas}", F_PCT)
    linea("r_mn", "Margen neto", "={col}{un}/{col}{vtas}", "={col}{un}/{col}{vtas}", "={col}{un}/{col}{vtas}", F_PCT)
    linea("r_rot", "Rotación de activos (ventas / activos)", "={col}{vtas}/{col}{ta}", "={col}{vtas}/{col}{ta}",
          "={col}{vtas}/{col}{ta}", F_X)
    linea("r_mult", "Multiplicador de capital (activos / patrimonio)", "={col}{ta}/{col}{pat}", "={col}{ta}/{col}{pat}",
          "={col}{ta}/{col}{pat}", F_X)
    linea("r_roe", "ROE (utilidad neta / patrimonio final)", "={col}{un}/{col}{pat}", "={col}{un}/{col}{pat}",
          "={col}{un}/{col}{pat}", F_PCT, negrita=True)
    linea("r_dup", "ROE por DuPont (margen × rotación × multiplicador)", "={col}{r_mn}*{col}{r_rot}*{col}{r_mult}",
          "={col}{r_mn}*{col}{r_rot}*{col}{r_mult}", "={col}{r_mn}*{col}{r_rot}*{col}{r_mult}", F_PCT)
    linea("r_roei", "ROE sobre patrimonio inicial (7.127)", None, "={col}{un}/$B{pat}", "={col}{un}/$B{pat}", F_PCT)
    linea("r_roa", "ROA (utilidad neta / activos)", "={col}{un}/{col}{ta}", "={col}{un}/{col}{ta}", "={col}{un}/{col}{ta}", F_PCT)
    linea("r_cob", "Cobertura de intereses (EBIT / intereses)", "={col}{ebit}/{col}{int}", "={col}{ebit}/{col}{int}",
          "={col}{ebit}/{col}{int}", F_X)
    linea("r_dcxc", "Días de cobro (CxC / ventas × 365)", "={col}{cxc}/{col}{vtas}*365", "={col}{cxc}/{col}{vtas}*365",
          "={col}{cxc}/{col}{vtas}*365", F_DIAS)
    linea("r_ctno", "Capital de trabajo operativo (inventarios + CxC – CxP)", "={col}{invmp}+{col}{invpt}+{col}{cxc}-{col}{cxp}",
          "={col}{invmp}+{col}{invpt}+{col}{cxc}-{col}{cxp}", "={col}{invmp}+{col}{invpt}+{col}{cxc}-{col}{cxp}")
    linea("r_ctnov", "Capital de trabajo operativo / ventas", "={col}{r_ctno}/{col}{vtas}", "={col}{r_ctno}/{col}{vtas}",
          "={col}{r_ctno}/{col}{vtas}", F_PCT)
    linea("r_rc", "Razón corriente", "={col}{ac}/{col}{pc}", "={col}{ac}/{col}{pc}", "={col}{ac}/{col}{pc}", F_X)
    linea("r_de", "Deuda financiera / patrimonio", "=({col}{lc}+{col}{dlp})/{col}{pat}", "=({col}{lc}+{col}{dlp})/{col}{pat}",
          "=({col}{lc}+{col}{dlp})/{col}{pat}", F_X)
    linea("r_dbt", "Deuda financiera / EBITDA", "=({col}{lc}+{col}{dlp})/{col}{ebitda}",
          "=({col}{lc}+{col}{dlp})/{col}{ebitda}", "=({col}{lc}+{col}{dlp})/{col}{ebitda}", F_X)
    linea("r_ebkg", "EBIT por kilo vendido (€/kg)", "={col}{ebit}/({col}{kg_mp}+{col}{kg_pr})*1000",
          "={col}{ebit}/({col}{kg_mp}+{col}{kg_pr})*1000", "={col}{ebit}/({col}{kg_mp}+{col}{kg_pr})*1000", F_PRECIO)
    linea("r_be", "Ventas de equilibrio (antes de intereses)", None,
          "=({col}{cfx}+{col}{dinc}+{col}{dbase}+{col}{gop}-{col}{mkt})/(1-({col}{cvmp}+{col}{cvpr})/{col}{vtas}-{col}{mkt}/{col}{vtas})",
          "=({col}{cfx}+{col}{dinc}+{col}{dbase}+{col}{gop}-{col}{mkt})/(1-({col}{cvmp}+{col}{cvpr})/{col}{vtas}-{col}{mkt}/{col}{vtas})")
    linea("r_bei", "Ventas de equilibrio (incluyendo intereses)", None,
          "=({col}{cfx}+{col}{dinc}+{col}{dbase}+{col}{gop}-{col}{mkt}+{col}{int})/(1-({col}{cvmp}+{col}{cvpr})/{col}{vtas}-{col}{mkt}/{col}{vtas})",
          "=({col}{cfx}+{col}{dinc}+{col}{dbase}+{col}{gop}-{col}{mkt}+{col}{int})/(1-({col}{cvmp}+{col}{cvpr})/{col}{vtas}-{col}{mkt}/{col}{vtas})")
    linea("r_ms", "Margen de seguridad (ventas sobre equilibrio con intereses)", None, "=1-{col}{r_bei}/{col}{vtas}",
          "=1-{col}{r_bei}/{col}{vtas}", F_PCT)

    for k, f in R.items():
        N["E_" + k] = f
    anchos(ws, {"A": 66, "B": 15, "C": 17, "D": 17, "E": 60})
    ws.freeze_panes = "B4"
    return R


# --------------------------------------------------------------------------
# Hoja Mejoras para marca propia
# --------------------------------------------------------------------------
def hoja_mejoras(wb, R):
    ws = wb.create_sheet("Mejoras_MP")
    S = lambda k: N[k]  # noqa: E731
    E = lambda k: f"EEFF_2020!$C${R[k]}"  # noqa: E731
    titulo(ws, "¿Se pueden mejorar los indicadores de la estrategia de solo marca propia? (miles €)",
           "Cada palanca se mide por separado contra el escenario base; la fila final las suma (aproximación: "
           "ignora interacciones menores).", 7)
    seccion(ws, 3, "Palancas (celdas amarillas editables)", 7)
    pal = [
        ("p_meses", "Meses de cobro negociados con los minoristas (base = 3)", 2, "0"),
        ("p_dprecio", "Aumento de precio negociado (€/kg)", 0.25, F_PRECIO),
        ("p_lote", "Venta del lote desocupado a valor libro, aplicada a reducir la línea de crédito", "=" + S("lote"), F_MILES),
        ("p_dpago", "Días de pago a proveedores (base ≈ 6 días)", 30, "0"),
        ("p_redmp", "Reducción del inventario de materia prima", 0.30, F_PCT),
    ]
    P = {}
    for i, (k, et, v, fmt) in enumerate(pal):
        f = 4 + i
        P[k] = f
        put(ws, f"A{f}", et)
        put(ws, f"B{f}", v, fmt, relleno=RELLENO_CLAVE)
    seccion(ws, 10, "Base: 2020 solo marca propia", 7)
    base = [("b_un", "Utilidad neta", E("un")), ("b_pat", "Patrimonio final", E("pat")),
            ("b_nec", "Excedente / (necesidad) de financiamiento", E("nec")),
            ("b_deuda", "Deuda financiera", f"{E('lc')}+{E('dlp')}"), ("b_cxc", "Cuentas a cobrar", E("cxc"))]
    B = {}
    for i, (k, et, fo) in enumerate(base):
        f = 11 + i
        B[k] = f
        put(ws, f"A{f}", et)
        put(ws, f"B{f}", "=" + fo, F_MILES)
    seccion(ws, 17, "Efecto de cada palanca", 7)
    encabezados(ws, 18, ["Palanca", "Δ Utilidad neta", "Δ Caja (+ libera / – consume)", "Δ Deuda", "Utilidad neta resultante",
                         "Excedente / (necesidad) de financiamiento resultante", "ROE resultante"])
    t = S("tasa_imp")
    v_mp = E("v_mp")
    filas = [
        ("Cobro en {meses} meses en lugar de 3",
         "0",
         f"=B{B['b_cxc']}-({v_mp}*SUM(OFFSET(Anexos!$B$36,1-$B${P['p_meses']},0,$B${P['p_meses']},1)))",
         "0"),
        ("Precio +Δ €/kg",
         f"=$B${P['p_dprecio']}*{E('kg_mp')}/1000*(1-{t})",
         f"=B20-$B${P['p_dprecio']}*{E('kg_mp')}/1000*{S('sh_cxc_mp')}",
         "0"),
        ("Venta del lote y reducción de deuda",
         f"=$B${P['p_lote']}*{S('tasa_int')}*(1-{t})",
         "=B21",
         f"=-$B${P['p_lote']}"),
        ("Alargar el pago a proveedores",
         "0",
         f"=({E('cmv')}-{E('dep')})*$B${P['p_dpago']}/365-{E('cxp')}",
         "0"),
        ("Reducir inventario de materia prima",
         "0",
         f"={E('invmp')}*$B${P['p_redmp']}",
         "0"),
    ]
    for i, (et, dun, dcaja, ddeu) in enumerate(filas):
        f = 19 + i
        put(ws, f"A{f}", et)
        put(ws, f"B{f}", dun if dun.startswith("=") else "=" + dun, F_MILES)
        put(ws, f"C{f}", dcaja, F_MILES)
        put(ws, f"D{f}", ddeu if ddeu.startswith("=") else "=" + ddeu, F_MILES)
        put(ws, f"E{f}", f"=$B${B['b_un']}+B{f}", F_MILES)
        put(ws, f"F{f}", f"=$B${B['b_nec']}+C{f}", F_MILES)
        put(ws, f"G{f}", f"=E{f}/($B${B['b_pat']}+B{f})", F_PCT)
    ws["A19"].value = '="Cobro en "&B4&" meses en lugar de 3"'
    ws["A20"].value = '="Precio +"&TEXT(B5,"0.00")&" €/kg"'
    f = 24
    put(ws, f"A{f}", "Todas las palancas juntas", negrita=True)
    put(ws, f"B{f}", "=SUM(B19:B23)", F_MILES, negrita=True)
    put(ws, f"C{f}", "=SUM(C19:C23)", F_MILES, negrita=True)
    put(ws, f"D{f}", "=SUM(D19:D23)", F_MILES, negrita=True)
    put(ws, f"E{f}", f"=$B${B['b_un']}+B{f}", F_MILES, negrita=True)
    put(ws, f"F{f}", f"=$B${B['b_nec']}+C{f}", F_MILES, negrita=True)
    put(ws, f"G{f}", f"=E{f}/($B${B['b_pat']}+B{f})", F_PCT, negrita=True)
    total(ws, 24, "ABCDEFG")
    put(ws, "A26", "Referencias", negrita=True)
    put(ws, "A27", "Utilidad neta 2019")
    put(ws, "B27", f"=EEFF_2020!$B${R['un']}", F_MILES)
    put(ws, "A28", "Utilidad neta 2020 – MIX óptimo")
    put(ws, "B28", f"=EEFF_2020!$D${R['un']}", F_MILES)
    put(ws, "A29", "Deuda / patrimonio con todas las palancas")
    put(ws, "B29", f"=($B${B['b_deuda']}+D24)/($B${B['b_pat']}+B24)", F_X)
    put(ws, "A31", "Lectura: las palancas de capital de trabajo reducen la necesidad de caja, pero la utilidad de la "
        "estrategia sigue muy por debajo de 2019 y del mix; el problema de fondo es el bajo margen unitario (1,75 €/kg).",
        cursiva=True)
    anchos(ws, {"A": 72, "B": 15, "C": 18, "D": 13, "E": 16, "F": 20, "G": 13})


# --------------------------------------------------------------------------
# Hoja Plan mensual
# --------------------------------------------------------------------------
def hoja_mensual(wb):
    ws = wb.create_sheet("Plan_Mensual")
    S = lambda k: N[k]  # noqa: E731
    titulo(ws, "Plan mensual 2020: ventas, producción, inventario y composición de las cuentas a cobrar",
           "Escenario solo marca propia replica el Anexo 5; en el mix la calidad superior no se almacena "
           "(se produce al ritmo de la venta) y la marca propia usa la capacidad restante.", 13)
    encabezados(ws, 4, ["Mes", "% ventas", "Capacidad (kg)",
                        "MP: ventas (kg)", "MP: producción (kg)", "MP: inventario final (kg)", "MP: ventas (miles €)",
                        "MIX: calidad sup. (kg)", "MIX: marca propia (kg)", "MIX: capacidad libre p/ MP (kg)",
                        "MIX: producción MP (kg)", "MIX: inventario MP (kg)", "MIX: ventas (miles €)"])
    put(ws, "A5", "Inventario inicial")
    put(ws, "F5", "=" + S("invpt_kg19"), F_KG)
    put(ws, "L5", "=" + S("invpt_kg19"), F_KG)
    for i in range(12):
        f = 6 + i
        put(ws, f"A{f}", MESES[i])
        put(ws, f"B{f}", f"=Anexos!B{25 + i}", F_PCT)
        put(ws, f"C{f}", "=" + S("cap_mes"), F_KG)
        # Solo marca propia (lógica del Anexo 5: inventario final de cada mes calculado desde diciembre hacia atrás)
        put(ws, f"D{f}", f"={S('cap')}*B{f}", F_KG)
        put(ws, f"E{f}", f"=MIN(C{f},MAX(0,D{f}+F{f}-F{f - 1}))", F_KG)
        put(ws, f"F{f}", f"=MAX(F{f + 1}+D{f + 1}-C{f + 1},0)" if i < 11 else "=" + S("invpt_kg20"), F_KG)
        put(ws, f"G{f}", f"=D{f}*{S('p_mp')}/1000", F_MILES)
        # Mix
        put(ws, f"H{f}", f"={S('opt_qpr')}*B{f}", F_KG)
        put(ws, f"I{f}", f"={S('opt_qmp')}*B{f}", F_KG)
        put(ws, f"J{f}", f"=C{f}-H{f}", F_KG)
        put(ws, f"K{f}", f"=MIN(J{f},MAX(0,I{f}+L{f}-L{f - 1}))", F_KG)
        put(ws, f"L{f}", f"=MAX(L{f + 1}+I{f + 1}-J{f + 1},0)" if i < 11 else "=" + S("invpt_kg20"), F_KG)
        put(ws, f"M{f}", f"=(H{f}*{S('opt_p')}+I{f}*{S('p_mp')})/1000", F_MILES)
    put(ws, "A18", "Total", negrita=True)
    for c in "BCDEGHIJKM":
        put(ws, f"{c}18", f"=SUM({c}6:{c}17)", F_PCT if c == "B" else (F_MILES if c in "GM" else F_KG), negrita=True)
    total(ws, 18, "ABCDEFGHIJKLM")
    put(ws, "A20", "Control MP: inicial + producción – ventas – final (debe ser 0)")
    put(ws, "B20", "=F5+E18-D18-F17", F_KG)
    put(ws, "A21", "Control MIX: inicial + producción MP – ventas MP – final (debe ser 0)")
    put(ws, "B21", "=L5+K18-I18-L17", F_KG)
    put(ws, "A22", "Control MIX: meses en que la calidad superior excede la capacidad (debe ser 0)")
    put(ws, "B22", "=COUNTIF(J6:J17,\"<0\")", "0")

    seccion(ws, 24, "Composición de las cuentas a cobrar al 31-dic-2020 (miles €)", 13)
    encabezados(ws, 25, ["Mes", "Solo MP", "MIX: calidad sup.", "MIX: marca propia", "MIX: total"])
    for i, mes in enumerate(["Octubre", "Noviembre", "Diciembre"]):
        f = 26 + i
        src = 15 + i  # filas de oct, nov, dic
        put(ws, f"A{f}", mes)
        put(ws, f"B{f}", f"=G{src}", F_MILES)
        put(ws, f"C{f}", f"=IF(\"{mes}\"=\"Diciembre\",H{src}*{S('opt_p')}/1000,0)", F_MILES)
        put(ws, f"D{f}", f"=I{src}*{S('p_mp')}/1000", F_MILES)
        put(ws, f"E{f}", f"=C{f}+D{f}", F_MILES)
    put(ws, "A29", "Total CxC", negrita=True)
    for c in "BCDE":
        put(ws, f"{c}29", f"=SUM({c}26:{c}28)", F_MILES, negrita=True)
    total(ws, 29, "ABCDE")
    put(ws, "A30", "Comparación: CxC al 31-dic-2019")
    put(ws, "B30", "=" + S("cxc19"), F_MILES)
    put(ws, "A31", "Nota: con 3 meses de cobro, la CxC solo de marca propia equivale al 37,2 % de las ventas anuales "
        "(oct 13,1 % + nov 12,1 % + dic 12,0 %), no a 90/365 = 24,7 %, porque las ventas no son estacionarias.",
        cursiva=True)
    anchos(ws, {"A": 20, **{get_column_letter(c): 14 for c in range(2, 14)}})
    ws.freeze_panes = "B5"


def main(salida):
    wb = Workbook()
    hoja_supuestos(wb)
    hoja_anexos(wb)
    hoja_diagnostico(wb)
    hoja_mix(wb)
    R = hoja_eeff(wb)
    hoja_mejoras(wb, R)
    hoja_mensual(wb)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(salida)


if __name__ == "__main__":
    main(sys.argv[1])
