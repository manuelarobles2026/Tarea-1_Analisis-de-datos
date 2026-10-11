"""Lee el Excel ya recalculado, exporta las cifras clave a JSON y genera los gráficos del informe.

Uso: python extraer_y_graficar.py <modelo.xlsx> <carpeta_salida>
"""
import json
import os
import sys

import matplotlib
import openpyxl

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

AZUL, NARANJA, AGUA = "#2a78d6", "#eb6834", "#1baf7a"
TINTA, TINTA2, GRILLA = "#0b0b0b", "#52514e", "#e4e3df"


def fila(ws, etiqueta, col_a=1):
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=col_a).value
        if isinstance(v, str) and v.strip() == etiqueta:
            return r
    raise KeyError(f"{ws.title}: no se encontró «{etiqueta}»")


def main(xlsx, salida):
    os.makedirs(salida, exist_ok=True)
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    d = {}

    ws = wb["EEFF_2020"]
    etiquetas = {
        "ventas": "Ventas totales", "v_mp": "Ventas marca propia", "v_pr": "Ventas calidad superior",
        "cmv": "Costo de mercadería vendida", "mb": "Margen bruto", "mkt": "Gastos de marketing",
        "id": "Gastos de I+D", "vta": "Gastos de venta", "adm": "Gastos administrativos",
        "gop": "Total gastos operativos", "ebit": "Utilidad operativa (EBIT)", "int": "Gastos de intereses",
        "uai": "Utilidad antes de impuestos", "imp": "Impuestos (40 %)", "un": "Utilidad neta",
        "dep": "Memo: depreciación total del año", "ebitda": "Memo: EBITDA",
        "caja": "Caja (variable de cierre)", "invmp": "Inventario de materia prima", "invpt": "Bienes terminados",
        "cxc": "Cuentas a cobrar", "ac": "Activo corriente", "ppe": "Bienes de uso (bruto)",
        "depac": "Depreciación acumulada", "ppen": "Bienes de uso netos", "ta": "Total activos",
        "cxp": "Cuentas a pagar", "lc": "Línea de crédito", "dlp": "Deuda a largo plazo", "tp": "Total pasivos",
        "pat": "Patrimonio neto", "tpp": "Total pasivos y patrimonio",
        "cajamin": "Memo: caja mínima operativa (% de ventas 2019)",
        "nec": "Memo: excedente / (necesidad) de financiamiento",
        "intad": "Memo: interés adicional si la necesidad se cubre con deuda (antes de impuestos)",
        "unaj": "Memo: utilidad neta ajustada por ese interés adicional",
        "f_dep": "(+) Depreciación", "f_cxc": "(–) Aumento de cuentas a cobrar",
        "f_imp": "(–) Aumento de inventario de materia prima", "f_ipt": "(–) Aumento de bienes terminados",
        "f_cxp": "(+) Aumento de cuentas a pagar", "fco": "Flujo de caja operativo",
        "fci": "Flujo de caja de inversión", "fcf": "Flujo de caja de financiamiento",
        "dcaja": "Variación neta de caja", "cini": "Caja inicial", "cfin": "Caja final",
        "r_mb": "Margen bruto", "r_ebit": "Margen operativo (EBIT / ventas)", "r_mn": "Margen neto",
        "r_rot": "Rotación de activos (ventas / activos)", "r_mult": "Multiplicador de capital (activos / patrimonio)",
        "r_roe": "ROE (utilidad neta / patrimonio final)", "r_roei": "ROE sobre patrimonio inicial (7.127)",
        "r_roa": "ROA (utilidad neta / activos)", "r_cob": "Cobertura de intereses (EBIT / intereses)",
        "r_dcxc": "Días de cobro (CxC / ventas × 365)",
        "r_ctno": "Capital de trabajo operativo (inventarios + CxC – CxP)",
        "r_ctnov": "Capital de trabajo operativo / ventas", "r_rc": "Razón corriente",
        "r_de": "Deuda financiera / patrimonio", "r_dbt": "Deuda financiera / EBITDA",
        "r_ebkg": "EBIT por kilo vendido (€/kg)", "r_bei": "Ventas de equilibrio (incluyendo intereses)",
        "r_ms": "Margen de seguridad (ventas sobre equilibrio con intereses)",
        "kg_mp": "Kilos marca propia", "kg_pr": "Kilos calidad superior",
    }
    # "Margen bruto" aparece dos veces (línea y ratio): el ratio es la segunda aparición
    filas_mb = [r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "Margen bruto"]
    for k, et in etiquetas.items():
        r = filas_mb[1] if k == "r_mb" else (filas_mb[0] if k == "mb" else fila(ws, et))
        d[k] = [ws.cell(row=r, column=c).value for c in (2, 3, 4)]

    ws = wb["Diagnostico_2019"]
    for k, et in {
        "d_int_ebit": "Intereses / EBIT", "d_rona": "RONA después de impuestos = EBIT(1–t) / (deuda + patrimonio)",
        "d_kd": "Costo de la deuda después de impuestos", "d_de": "Deuda financiera / patrimonio",
        "d_dcxc": "Días de cobro (CxC / ventas × 365)", "d_dinv": "Días de inventario MP (sobre CMV sin depreciación)",
        "d_dpt": "Días de bienes terminados", "d_dpag": "Días de pago (CxP / CMV sin depreciación × 365)",
        "d_ccc": "Ciclo de conversión de efectivo (días)", "d_ctno": "Capital de trabajo operativo (inventarios + CxC – CxP)",
        "d_ctnov": "Capital de trabajo operativo / ventas", "d_rc": "Razón corriente (activo corriente / pasivo corriente)",
        "d_acida": "Prueba ácida ((caja + CxC) / pasivo corriente)", "d_lote": "Lote desocupado / total activos",
        "d_cxc_est": "CxC estimada total vs. CxC del balance", "d_p_mp": "Precio promedio marca propia (€/kg)",
        "d_p_pr": "Precio promedio calidad superior (€/kg)", "d_uso19": "Uso de la capacidad 2019 (350.000 kg/mes)",
        "d_uso20": "Uso de la capacidad 2020 con el mismo volumen", "d_dbt": "Deuda total / EBITDA",
    }.items():
        d[k] = ws.cell(row=fila(ws, et), column=2).value

    ws = wb["Mix_Produccion"]
    r_cal = fila(ws, "Calidad")
    # La primera aparición de «Calidad» está en la Tabla A; el resumen usa la fila de «Mejor resultado del MIX»
    r_res = fila(ws, "Mejor resultado del MIX (miles €)")
    r_cal = r_res - 1
    d["mix_grados"] = [ws.cell(row=r_cal, column=c).value for c in range(3, 9)]
    for k, et in {
        "mix_best": "Mejor resultado del MIX (miles €)", "mix_m": "% marketing óptimo en el MIX",
        "mix_qpr": "Kilos de calidad superior", "mix_qmp": "Kilos de marca propia",
        "solo_best": "Mejor resultado SOLO calidad superior (miles €)", "solo_m": "% marketing óptimo SOLO calidad superior",
        "solo_q": "Kilos SOLO calidad superior", "solo_ocio": "Capacidad ociosa SOLO calidad superior",
    }.items():
        r = fila(ws, et)
        d[k] = [ws.cell(row=r, column=c).value for c in range(3, 9)]
    d["mp_solo_mini"] = ws.cell(row=fila(ws, "Resultado solo marca propia (D) a plena capacidad"), column=2).value
    for k, et in {"opt_res": "Resultado máximo (miles €)", "opt_cal": "Calidad de alta gama elegida",
                  "opt_m": "% de marketing / ventas", "opt_peso": "Peso de calidad superior en el volumen"}.items():
        d[k] = ws.cell(row=fila(ws, et), column=2).value
    mini = {}
    r0 = fila(ws, "Calidad superior: kilos")
    for r in range(r0, r0 + 19):
        et = ws.cell(row=r, column=1).value
        if et:
            mini[et] = [ws.cell(row=r, column=c).value for c in range(2, 6)]
    d["mini"] = mini
    # Tabla C completa (para el anexo del informe)
    r_tc = fila(ws, "Tabla C – Resultado del MIX: calidad superior + marca propia en la capacidad restante (miles €)")
    d["tabla_c"] = [[ws.cell(row=r, column=c).value for c in range(1, 9)] for r in range(r_tc + 5, r_tc + 16)]

    ws = wb["Mejoras_MP"]
    d["mejoras"] = [[ws.cell(row=r, column=c).value for c in range(1, 8)] for r in range(19, 25)]

    with open(os.path.join(salida, "datos.json"), "w", encoding="utf-8") as fh:
        json.dump(d, fh, ensure_ascii=False, indent=1)

    # ---------------- Gráficos ----------------
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "axes.edgecolor": GRILLA,
                         "axes.labelcolor": TINTA2, "xtick.color": TINTA2, "ytick.color": TINTA2})

    def miles(x):
        return f"{x:,.0f}".replace(",", ".")

    # 1) Mejor resultado por calidad
    fig, ax = plt.subplots(figsize=(7.5, 3.9), dpi=200)
    g = d["mix_grados"]
    xs = range(len(g))
    w = 0.38
    b1 = ax.bar([x - w / 2 - 0.01 for x in xs], d["mix_best"], w, color=AZUL, label="Mix: calidad superior + marca propia")
    b2 = ax.bar([x + w / 2 + 0.01 for x in xs], d["solo_best"], w, color=NARANJA, label="Solo calidad superior")
    ax.axhline(d["mp_solo_mini"], color=TINTA2, lw=1.2, ls="--",
               label=f"Solo marca propia ({miles(d['mp_solo_mini'])})")
    i_opt = g.index(d["opt_cal"])
    ax.text(i_opt - w / 2, d["mix_best"][i_opt] + 400, miles(d["mix_best"][i_opt]), ha="center", va="bottom",
            fontsize=9, color=TINTA, fontweight="bold")
    i_aaa = g.index("AAA")
    ax.text(i_aaa + w / 2, d["solo_best"][i_aaa] + 400, miles(d["solo_best"][i_aaa]), ha="center", va="bottom",
            fontsize=8.5, color=TINTA)
    ax.set_xticks(list(xs))
    ax.set_xticklabels([f"{c}\n%M {d['mix_m'][i] * 100:.0f} %" for i, c in enumerate(g)], fontsize=8.5)
    ax.set_ylabel("Miles de €")
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: miles(v)))
    ax.grid(axis="y", color=GRILLA, lw=0.8)
    ax.set_axisbelow(True)
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.set_ylim(0, max(d["mix_best"]) * 1.25)
    ax.legend(frameon=False, fontsize=8, loc="upper left", ncol=3)
    ax.set_xlabel("Calidad de alta gama y % de marketing óptimo en el mix", fontsize=8.5)
    ax.set_title("Mejor resultado (p – cv)·q – CF – %M·p·q por calidad", fontsize=11, color=TINTA, loc="left")
    del b1, b2
    fig.tight_layout()
    fig.savefig(os.path.join(salida, "grafico_mix.png"))
    plt.close(fig)

    # 2) Utilidad neta 2020 por alternativa
    m = d["mini"]["Utilidad neta (40 % de impuestos)"]
    nombres = ["2019 real", "Solo marca\npropia", "Solo AAA", f"Solo calidad {d['opt_cal']}", "Mix óptimo"]
    valores = [d["un"][0], m[0], m[1], m[2], m[3]]
    fig, ax = plt.subplots(figsize=(7.5, 3.4), dpi=200)
    ax.bar(range(5), valores, 0.55, color=AZUL)
    for i, v in enumerate(valores):
        ax.text(i, v + (250 if v >= 0 else -250), miles(v), ha="center", va="bottom" if v >= 0 else "top",
                fontsize=9, color=TINTA)
    ax.axhline(0, color=TINTA2, lw=0.8)
    ax.set_xticks(range(5))
    ax.set_xticklabels(nombres, fontsize=9)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: miles(v)))
    ax.grid(axis="y", color=GRILLA, lw=0.8)
    ax.set_axisbelow(True)
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.set_ylim(min(valores) * 1.5, max(valores) * 1.18)
    ax.set_ylabel("Miles de €")
    ax.set_title("Utilidad neta 2020 por alternativa (miles de €)", fontsize=11, color=TINTA, loc="left")
    fig.tight_layout()
    fig.savefig(os.path.join(salida, "grafico_utilidad.png"))
    plt.close(fig)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
